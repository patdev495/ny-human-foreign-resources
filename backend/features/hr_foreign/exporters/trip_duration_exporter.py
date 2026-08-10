from __future__ import annotations

import datetime
import io
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import or_
from sqlalchemy.orm import Session

from features.hr_foreign.models import ForeignEmployee, Stay, TravelRecord
from features.hr_foreign.status_engine import evaluate_employee_statuses
from .excel_styles import (
    BOLD_FONT,
    BORDER_THIN,
    HEADER_FONT,
    REGULAR_FONT,
    TITLE_FONT,
    apply_header_style,
    auto_fit_columns,
    format_work_type,
)


def generate_trip_duration_excel(
    db: Session,
    start_date: datetime.date,
    end_date: datetime.date,
    employee_id: int | None = None,
    today: datetime.date | None = None,
) -> io.BytesIO:
    if today is None:
        today = datetime.date.today()

    query = (
        db.query(TravelRecord)
        .join(ForeignEmployee)
        .filter(
            or_(
                ForeignEmployee.employee_type != "JANITORIAL",
                ForeignEmployee.employee_type.is_(None),
            )
        )
    )

    if employee_id is not None:
        query = query.filter(TravelRecord.employee_id == employee_id)

    # Filter records that overlap with [start_date, end_date]
    # Entry date <= end_date AND (actual_exit_date is NULL OR actual_exit_date >= start_date)
    records = (
        query.filter(
            or_(TravelRecord.entry_date.is_(None), TravelRecord.entry_date <= end_date),
            or_(TravelRecord.actual_exit_date.is_(None), TravelRecord.actual_exit_date >= start_date),
        )
        .order_by(ForeignEmployee.name_latin, TravelRecord.entry_date)
        .all()
    )

    # Pre-evaluate employee statuses for current accommodation info
    all_employees = (
        db.query(ForeignEmployee)
        .filter(
            or_(
                ForeignEmployee.employee_type != "JANITORIAL",
                ForeignEmployee.employee_type.is_(None),
            )
        )
        .all()
    )
    evaluated_list = evaluate_employee_statuses(db, all_employees, today=today)
    eval_map = {e.id: e for e in evaluated_list}

    wb = Workbook()
    ws = wb.active
    ws.title = "Đợt Lưu Trú & Nhập Xuất Cảnh"

    headers = [
        "STT",
        "Mã NV",
        "Họ tên Latin",
        "Họ tên Trung Quốc",
        "Số Hộ chiếu",
        "Bộ phận / Chức danh",
        "Loại hình",
        "Chỗ ở hiện tại",
        "Ngày Đến VN",
        "Ngày Về nước thực tế",
        "Ngày Dự kiến về",
        "Số ngày ở trong kỳ",
        "Tổng số ngày đợt di chuyển",
        "Ghi chú",
    ]
    apply_header_style(ws, headers)

    row_count = 0
    for idx, tr in enumerate(records, start=1):
        row_count += 1
        emp = tr.employee
        emp_eval = eval_map.get(tr.employee_id)

        room_info = "–"
        if emp_eval and emp_eval.is_in_vietnam:
            room_info = emp_eval.current_room_number or "Khách sạn"

        entry_d = tr.entry_date
        exit_d = tr.actual_exit_date

        # Calculate overlap with [start_date, end_date]
        trip_start = entry_d or start_date
        trip_end = exit_d or today

        overlap_start = max(trip_start, start_date)
        overlap_end = min(trip_end, end_date)

        if overlap_end >= overlap_start:
            days_in_period = (overlap_end - overlap_start).days + 1
        else:
            days_in_period = 0

        if trip_end >= trip_start:
            total_trip_days = (trip_end - trip_start).days + 1
        else:
            total_trip_days = 0

        dept_role = emp.department or ""
        if emp.role:
            dept_role += f" ({emp.role})" if dept_role else emp.role

        ws.append([
            idx,
            emp.employee_code or "",
            emp.name_latin or "",
            emp.name_chinese or "",
            emp.passport_number or "",
            dept_role,
            format_work_type(emp.work_type),
            room_info,
            entry_d.strftime("%d/%m/%Y") if entry_d else "",
            exit_d.strftime("%d/%m/%Y") if exit_d else ("Đang ở VN" if not exit_d else ""),
            tr.expected_exit_date.strftime("%d/%m/%Y") if tr.expected_exit_date else "",
            days_in_period,
            total_trip_days,
            tr.notes or "",
        ])

        curr_r = idx + 1
        for c in range(1, 15):
            cell = ws.cell(row=curr_r, column=c)
            cell.font = REGULAR_FONT
            cell.border = BORDER_THIN
            if c in [1, 7, 9, 10, 11]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif c in [12, 13]:
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = "#,##0"

    # Summary Row at the bottom
    tot_row = row_count + 2
    cell_label = ws.cell(row=tot_row, column=1, value="TỔNG CỘNG")
    cell_label.font = BOLD_FONT
    cell_label.border = BORDER_THIN
    cell_label.alignment = Alignment(horizontal="center", vertical="center")

    for c in range(2, 15):
        cell = ws.cell(row=tot_row, column=c)
        cell.font = BOLD_FONT
        cell.border = BORDER_THIN
        if c == 12:
            if row_count > 0:
                cell.value = f"=SUM(L2:L{row_count + 1})"
            else:
                cell.value = 0
            cell.alignment = Alignment(horizontal="right", vertical="center")
            cell.number_format = "#,##0"
        elif c == 13:
            if row_count > 0:
                cell.value = f"=SUM(M2:M{row_count + 1})"
            else:
                cell.value = 0
            cell.alignment = Alignment(horizontal="right", vertical="center")
            cell.number_format = "#,##0"
        else:
            cell.alignment = Alignment(horizontal="center", vertical="center")


    auto_fit_columns(ws)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
