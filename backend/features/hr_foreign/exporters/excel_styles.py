from __future__ import annotations

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

NAVY_FILL = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
BOLD_FONT = Font(name="Calibri", size=11, bold=True)
TITLE_FONT = Font(name="Calibri", size=14, bold=True)
RED_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
RED_FONT = Font(name="Calibri", size=11, color="9C0006", bold=True)
YELLOW_FILL = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
YELLOW_FONT = Font(name="Calibri", size=11, color="9C6500", bold=True)
HYPERLINK_FONT = Font(name="Calibri", size=11, color="0000FF", underline="single")
REGULAR_FONT = Font(name="Calibri", size=11)

BORDER_THIN = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)

WORK_TYPE_MAP = {
    "CO_DINH": "Cố định",
    "CONG_TAC": "Công tác",
    "Cố định": "Cố định",
    "Công tác": "Công tác",
}


def format_work_type(val: str | None) -> str:
    if not val:
        return "Cố định"
    return WORK_TYPE_MAP.get(val, val)


def build_attachment_zip_folder(emp_code: str | None, emp_name: str | None) -> str:
    code = emp_code or "NV_KHONG_MA"
    name = emp_name or "NHAN_VIEN"
    return f"Giay_To_Dinh_Kem/{code}_{name}".replace(" ", "_")


def build_attachment_zip_path(emp_code: str | None, emp_name: str | None, file_name: str) -> str:
    folder_name = build_attachment_zip_folder(emp_code, emp_name)
    return f"{folder_name}/{file_name}"


def apply_header_style(ws, headers: list[str]) -> None:
    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = NAVY_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"


def auto_fit_columns(ws) -> None:
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
