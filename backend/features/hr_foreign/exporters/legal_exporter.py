from __future__ import annotations

import datetime
import io
from openpyxl import Workbook
from sqlalchemy import or_
from sqlalchemy.orm import Session

from features.hr_foreign.models import DocumentAttachment, ForeignEmployee
from features.hr_foreign.status_engine import evaluate_employee_statuses
from .excel_styles import (
    HYPERLINK_FONT,
    RED_FILL,
    RED_FONT,
    YELLOW_FILL,
    YELLOW_FONT,
    apply_header_style,
    auto_fit_columns,
    build_attachment_zip_folder,
    build_attachment_zip_path,
    format_work_type,
)


def generate_legal_profile_excel(db: Session, today: datetime.date | None = None) -> io.BytesIO:
    if today is None:
        today = datetime.date.today()

    employees_db = (
        db.query(ForeignEmployee)
        .filter(
            or_(
                ForeignEmployee.employee_type != "JANITORIAL",
                ForeignEmployee.employee_type.is_(None),
            )
        )
        .order_by(ForeignEmployee.name_latin)
        .all()
    )
    evaluated = evaluate_employee_statuses(db, employees_db, today=today)

    all_attachments = db.query(DocumentAttachment).all()
    attachment_map: dict[tuple[str, int], list[DocumentAttachment]] = {}
    for att in all_attachments:
        key = (att.entity_type, att.entity_id)
        if key not in attachment_map:
            attachment_map[key] = []
        attachment_map[key].append(att)

    def _set_hyperlink_with_atts(cell, emp_code: str | None, emp_name: str | None, atts: list[DocumentAttachment] | None) -> bool:
        if not atts:
            return False
        valid_atts = [a for a in atts if a.file_name]
        if not valid_atts:
            return False

        cell.font = HYPERLINK_FONT
        if len(valid_atts) == 1:
            cell.hyperlink = build_attachment_zip_path(emp_code, emp_name, valid_atts[0].file_name)
            if not cell.value or str(cell.value).strip() == "":
                cell.value = "Xem file"
        else:
            cell.hyperlink = build_attachment_zip_folder(emp_code, emp_name)
            if not cell.value or str(cell.value).strip() == "":
                cell.value = f"Xem thư mục ({len(valid_atts)} file)"
        return True

    def _set_hyperlink_if_exists(cell, emp_code: str | None, emp_name: str | None, entity_type: str, entity_id: int | None) -> bool:
        if not entity_id:
            return False
        atts = attachment_map.get((entity_type, entity_id))
        return _set_hyperlink_with_atts(cell, emp_code, emp_name, atts)

    wb = Workbook()
    
    # Sheet 1: Danh sách Nhân sự
    ws1 = wb.active
    ws1.title = "Danh sách Nhân sự"
    headers1 = [
        "STT", "Mã NV", "Họ tên Latin", "Họ tên Trung Quốc", "Giới tính", "Ngày sinh", "Quốc tịch",
        "Số điện thoại", "Vị trí/Bộ phận", "Chức danh", "Số Hộ chiếu", "Hạn Hộ chiếu",
        "Ngày Đến VN", "Ngày Dự kiến sang", "Ngày Dự kiến về", "Ngày Thực tế về", "Hiện diện VN",
        "Chỗ ở hiện tại", "Loại hình", "Đăng ký ăn", "Loại Visa", "Hạn Visa", "Số ngày còn Visa",
        "Hạn Tạm trú", "Số ngày còn Tạm trú", "Số GPLĐ", "Cấp GPLĐ", "Hạn GPLĐ", "Số ngày còn GPLĐ",
        "Số HĐLĐ", "Loại HĐLĐ", "Hạn HĐLĐ", "Số ngày còn HĐLĐ", "Ghi chú",
    ]
    apply_header_style(ws1, headers1)

    for idx, emp_eval in enumerate(evaluated, start=1):
        emp_obj = next((e for e in employees_db if e.id == emp_eval.id), None)

        def days_rem(dt: datetime.date | None) -> int | None:
            if dt is None:
                return None
            return (dt - today).days

        visa_rem = days_rem(emp_eval.latest_visa_expiry)
        tamtru_rem = days_rem(emp_eval.latest_tamtru_expiry)
        gpld_rem = days_rem(emp_eval.latest_gpld_expiry)
        contract_rem = days_rem(emp_eval.latest_contract_expiry)

        latest_wp = None
        if emp_obj and emp_obj.work_permits:
            sorted_wps = sorted(emp_obj.work_permits, key=lambda w: w.valid_to or datetime.date.min, reverse=True)
            latest_wp = sorted_wps[0]

        latest_contract = None
        if emp_obj and emp_obj.contracts:
            sorted_contracts = sorted(emp_obj.contracts, key=lambda c: c.end_date or datetime.date.min, reverse=True)
            latest_contract = sorted_contracts[0]

        row_data = [
            idx,
            emp_eval.employee_code or "",
            emp_eval.name_latin or "",
            emp_eval.name_chinese or "",
            emp_eval.gender or "",
            emp_eval.date_of_birth.strftime("%d/%m/%Y") if emp_eval.date_of_birth else "",
            emp_eval.nationality or "",
            emp_eval.phone or "",
            emp_eval.department or "",
            emp_eval.role or "",
            emp_eval.passport_number or "",
            emp_eval.passport_expiry.strftime("%d/%m/%Y") if emp_eval.passport_expiry else "",
            emp_eval.entry_date.strftime("%d/%m/%Y") if emp_eval.entry_date else "",
            emp_eval.expected_entry_date.strftime("%d/%m/%Y") if emp_eval.expected_entry_date else "",
            emp_eval.expected_exit_date.strftime("%d/%m/%Y") if emp_eval.expected_exit_date else "",
            emp_eval.actual_exit_date.strftime("%d/%m/%Y") if emp_eval.actual_exit_date else "",
            "Đang ở VN" if emp_eval.is_in_vietnam else "Đã về nước",
            emp_eval.current_room_number or ("Khách sạn" if emp_eval.is_in_vietnam else "–"),
            format_work_type(emp_eval.work_type),
            "Có" if emp_eval.is_in_vietnam else "Không",
            emp_eval.latest_visa_type or "",
            emp_eval.latest_visa_expiry.strftime("%d/%m/%Y") if emp_eval.latest_visa_expiry else "",
            visa_rem if visa_rem is not None else "",
            emp_eval.latest_tamtru_expiry.strftime("%d/%m/%Y") if emp_eval.latest_tamtru_expiry else "",
            tamtru_rem if tamtru_rem is not None else "",
            latest_wp.permit_number if (latest_wp and latest_wp.permit_number) else "",
            latest_wp.issue_type if (latest_wp and latest_wp.issue_type) else "",
            emp_eval.latest_gpld_expiry.strftime("%d/%m/%Y") if emp_eval.latest_gpld_expiry else "",
            gpld_rem if gpld_rem is not None else "",
            latest_contract.contract_number if (latest_contract and latest_contract.contract_number) else "",
            latest_contract.contract_type if (latest_contract and latest_contract.contract_type) else "",
            emp_eval.latest_contract_expiry.strftime("%d/%m/%Y") if emp_eval.latest_contract_expiry else "",
            contract_rem if contract_rem is not None else "",
            emp_eval.notes or "",
        ]
        ws1.append(row_data)

        curr_row = idx + 1

        def _get_atts_for_ids(entity_type: str, entity_ids: list[int]) -> list[DocumentAttachment]:
            res: list[DocumentAttachment] = []
            for eid in entity_ids:
                a_list = attachment_map.get((entity_type, eid))
                if a_list:
                    res.extend(a_list)
            return res

        if emp_obj:
            passport_atts = attachment_map.get(("PASSPORT", emp_obj.id))
            _set_hyperlink_with_atts(ws1.cell(row=curr_row, column=11), emp_eval.employee_code, emp_eval.name_latin, passport_atts)

        if emp_obj and emp_obj.stays:
            visa_ids = [v.id for s in emp_obj.stays for v in s.visas]
            visa_atts = _get_atts_for_ids("VISA", visa_ids)
            _set_hyperlink_with_atts(ws1.cell(row=curr_row, column=21), emp_eval.employee_code, emp_eval.name_latin, visa_atts)

            tt_ids = [tt.id for s in emp_obj.stays for tt in s.tam_trus]
            tt_atts = _get_atts_for_ids("TAM_TRU", tt_ids)
            _set_hyperlink_with_atts(ws1.cell(row=curr_row, column=24), emp_eval.employee_code, emp_eval.name_latin, tt_atts)

        if emp_obj and emp_obj.work_permits:
            wp_ids = [wp.id for wp in emp_obj.work_permits]
            wp_atts = _get_atts_for_ids("WORK_PERMIT", wp_ids)
            _set_hyperlink_with_atts(ws1.cell(row=curr_row, column=26), emp_eval.employee_code, emp_eval.name_latin, wp_atts)

        if emp_obj and emp_obj.contracts:
            ct_ids = [c.id for c in emp_obj.contracts]
            ct_atts = _get_atts_for_ids("CONTRACT", ct_ids)
            _set_hyperlink_with_atts(ws1.cell(row=curr_row, column=30), emp_eval.employee_code, emp_eval.name_latin, ct_atts)

        def style_rem(col_idx: int, days: int | None):
            if days is None:
                return
            cell = ws1.cell(row=curr_row, column=col_idx)
            if days < 0:
                cell.fill = RED_FILL
                cell.font = RED_FONT
            elif days <= 30:
                cell.fill = YELLOW_FILL
                cell.font = YELLOW_FONT

        style_rem(23, visa_rem)
        style_rem(25, tamtru_rem)
        style_rem(29, gpld_rem)
        style_rem(33, contract_rem)

    auto_fit_columns(ws1)

    # Sheet 2: Lịch sử Giấy tờ
    ws2 = wb.create_sheet(title="Lịch sử Giấy tờ")
    headers2 = ["STT", "Mã NV", "Họ tên Latin", "Loại Giấy tờ", "Số / Phân loại", "Ngày Cấp / Đăng ký", "Ngày Hết hạn", "Ghi chú"]
    apply_header_style(ws2, headers2)

    s2_idx = 1
    for emp in employees_db:
        for stay in emp.stays:
            for v in stay.visas:
                s2_row = s2_idx + 1
                ws2.append([
                    s2_idx, emp.employee_code or "", emp.name_latin, "VISA",
                    v.visa_type or "", v.entry_date.strftime("%d/%m/%Y") if v.entry_date else "",
                    v.expiry_date.strftime("%d/%m/%Y") if v.expiry_date else "", v.notes or ""
                ])
                _set_hyperlink_if_exists(ws2.cell(row=s2_row, column=5), emp.employee_code, emp.name_latin, "VISA", v.id)
                s2_idx += 1

            for tt in stay.tam_trus:
                s2_row = s2_idx + 1
                ws2.append([
                    s2_idx, emp.employee_code or "", emp.name_latin, "TAM_TRU",
                    "Đăng ký Tạm trú", tt.registration_date.strftime("%d/%m/%Y") if tt.registration_date else "",
                    tt.expiry_date.strftime("%d/%m/%Y") if tt.expiry_date else "", tt.notes or ""
                ])
                _set_hyperlink_if_exists(ws2.cell(row=s2_row, column=5), emp.employee_code, emp.name_latin, "TAM_TRU", tt.id)
                s2_idx += 1

        for wp in emp.work_permits:
            s2_row = s2_idx + 1
            ws2.append([
                s2_idx, emp.employee_code or "", emp.name_latin, "GPLD",
                wp.permit_number or wp.issue_type or "", wp.valid_from.strftime("%d/%m/%Y") if wp.valid_from else "",
                wp.valid_to.strftime("%d/%m/%Y") if wp.valid_to else "", wp.notes or ""
            ])
            _set_hyperlink_if_exists(ws2.cell(row=s2_row, column=5), emp.employee_code, emp.name_latin, "WORK_PERMIT", wp.id)
            s2_idx += 1

        for c in emp.contracts:
            s2_row = s2_idx + 1
            ws2.append([
                s2_idx, emp.employee_code or "", emp.name_latin, "CONTRACT",
                c.contract_number or c.contract_type or "", c.start_date.strftime("%d/%m/%Y") if c.start_date else "",
                c.end_date.strftime("%d/%m/%Y") if c.end_date else "", c.notes or ""
            ])
            _set_hyperlink_if_exists(ws2.cell(row=s2_row, column=5), emp.employee_code, emp.name_latin, "CONTRACT", c.id)
            s2_idx += 1

    auto_fit_columns(ws2)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
