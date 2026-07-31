from __future__ import annotations

import datetime
import io
from openpyxl import Workbook
from sqlalchemy.orm import Session

from features.hr_foreign.models import ForeignEmployee
from features.hr_foreign.status_engine import evaluate_employee_statuses
from .excel_styles import apply_header_style, auto_fit_columns, format_work_type


def generate_presence_accommodation_excel(db: Session, today: datetime.date | None = None) -> io.BytesIO:
    if today is None:
        today = datetime.date.today()

    employees_db = db.query(ForeignEmployee).order_by(ForeignEmployee.name_latin).all()
    evaluated = evaluate_employee_statuses(db, employees_db, today=today)

    wb = Workbook()

    # Sheet 1: Sơ đồ KTX & Khách sạn
    ws1 = wb.active
    ws1.title = "Sơ đồ KTX & Khách sạn"
    headers1 = [
        "STT", "Loại chỗ ở", "Tên Cơ sở / Phòng", "Số phòng KS", "Vị trí giường",
        "Mã NV", "Họ tên Latin", "Họ tên Trung Quốc", "Số Hộ chiếu", "Bộ phận",
        "Loại hình", "Đăng ký ăn", "Số tiền hóa đơn đợt ở", "Ngày vào ở", "Ngày dự kiến về"
    ]
    apply_header_style(ws1, headers1)

    active_assigned = [e for e in evaluated if e.is_in_vietnam and e.current_room_number]
    for idx, emp_eval in enumerate(active_assigned, start=1):
        emp_obj = next((e for e in employees_db if e.id == emp_eval.id), None)
        active_stay = None
        if emp_obj:
            active_stay = next((s for s in emp_obj.stays if s.end_date is None or s.end_date > today), None)

        acc_type = active_stay.accommodation_type if active_stay else "KTX"
        unit_name = emp_eval.current_room_number or ""
        bed_loc = active_stay.bed_location if active_stay else ""
        hotel_room = active_stay.hotel_room_number if active_stay else ""
        stay_type = active_stay.stay_type if active_stay else "CO_DINH"
        has_meals = "Có" if (active_stay and active_stay.has_meals) else "Không"
        inv_amt = active_stay.invoice_amount if (active_stay and active_stay.invoice_amount is not None) else ""
        start_date = active_stay.start_date.strftime("%d/%m/%Y") if (active_stay and active_stay.start_date) else ""
        exp_end = emp_eval.expected_exit_date.strftime("%d/%m/%Y") if emp_eval.expected_exit_date else ""

        ws1.append([
            idx, acc_type, unit_name, hotel_room, bed_loc,
            emp_eval.employee_code or "", emp_eval.name_latin, emp_eval.name_chinese or "",
            emp_eval.passport_number or "", emp_eval.department or "",
            format_work_type(stay_type), has_meals, inv_amt, start_date, exp_end
        ])
    auto_fit_columns(ws1)

    # Sheet 2: Chưa xếp chỗ ở
    ws2 = wb.create_sheet(title="Chưa xếp chỗ ở")
    headers2 = ["STT", "Mã NV", "Họ tên Latin", "Họ tên Trung Quốc", "Số Hộ chiếu", "Bộ phận", "Ngày đến VN", "Ngày dự kiến về"]
    apply_header_style(ws2, headers2)

    unassigned = [e for e in evaluated if e.is_in_vietnam and not e.current_room_number]
    for idx, emp_eval in enumerate(unassigned, start=1):
        ws2.append([
            idx, emp_eval.employee_code or "", emp_eval.name_latin, emp_eval.name_chinese or "",
            emp_eval.passport_number or "", emp_eval.department or "",
            emp_eval.entry_date.strftime("%d/%m/%Y") if emp_eval.entry_date else "",
            emp_eval.expected_exit_date.strftime("%d/%m/%Y") if emp_eval.expected_exit_date else ""
        ])
    auto_fit_columns(ws2)

    # Sheet 3: Đã về nước
    ws3 = wb.create_sheet(title="Đã về nước")
    headers3 = ["STT", "Mã NV", "Họ tên Latin", "Họ tên Trung Quốc", "Số Hộ chiếu", "Bộ phận", "Ngày thực tế đã về", "Ngày dự kiến sang đợt tới"]
    apply_header_style(ws3, headers3)

    exited = [e for e in evaluated if not e.is_in_vietnam]
    for idx, emp_eval in enumerate(exited, start=1):
        ws3.append([
            idx, emp_eval.employee_code or "", emp_eval.name_latin, emp_eval.name_chinese or "",
            emp_eval.passport_number or "", emp_eval.department or "",
            emp_eval.actual_exit_date.strftime("%d/%m/%Y") if emp_eval.actual_exit_date else "",
            emp_eval.expected_entry_date.strftime("%d/%m/%Y") if emp_eval.expected_entry_date else ""
        ])
    auto_fit_columns(ws3)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
