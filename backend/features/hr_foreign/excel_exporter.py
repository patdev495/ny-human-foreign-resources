from __future__ import annotations

import io
import os
import zipfile
import datetime
from typing import Any
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session, joinedload

from features.hr_foreign.models import (
    ForeignEmployee,
    Stay,
    Room,
    Hotel,
    Visa,
    TamTru,
    WorkPermit,
    Contract,
    DocumentAttachment,
)
from features.hr_foreign.status_engine import evaluate_employee_statuses


NAVY_FILL = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
RED_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
RED_FONT = Font(name="Calibri", size=11, color="9C0006", bold=True)
YELLOW_FILL = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
YELLOW_FONT = Font(name="Calibri", size=11, color="9C6500", bold=True)
HYPERLINK_FONT = Font(name="Calibri", size=11, color="0000FF", underline="single")
REGULAR_FONT = Font(name="Calibri", size=11)
BORDER_THIN = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)


WORK_TYPE_MAP = {
    "CO_DINH": "Cố định",
    "CONG_TAC": "Công tác",
    "Cố định": "Cố định",
    "Công tác": "Công tác",
}


def format_work_type(val: str | None) -> str:
    if not val:
        return "Cố định"
    return WORK_TYPE_MAP.get(val, val)


def build_attachment_zip_folder(emp_code: str | None, emp_name: str | None) -> str:
    code = emp_code or "NV_KHONG_MA"
    name = emp_name or "NHAN_VIEN"
    return f"Giay_To_Dinh_Kem/{code}_{name}".replace(" ", "_")


def build_attachment_zip_path(emp_code: str | None, emp_name: str | None, file_name: str) -> str:
    folder_name = build_attachment_zip_folder(emp_code, emp_name)
    return f"{folder_name}/{file_name}"



def _apply_header_style(ws, headers: list[str]) -> None:
    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = NAVY_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"


def _auto_fit_columns(ws) -> None:
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)


def generate_legal_profile_excel(db: Session, today: datetime.date | None = None) -> io.BytesIO:
    if today is None:
        today = datetime.date.today()

    employees_db = db.query(ForeignEmployee).order_by(ForeignEmployee.name_latin).all()
    evaluated = evaluate_employee_statuses(db, employees_db, today=today)

    # Pre-fetch all document attachments and map by (entity_type, entity_id)
    all_attachments = db.query(DocumentAttachment).all()
    attachment_map: dict[tuple[str, int], list[DocumentAttachment]] = {}
    for att in all_attachments:
        key = (att.entity_type, att.entity_id)
        if key not in attachment_map:
            attachment_map[key] = []
        attachment_map[key].append(att)

    def _set_hyperlink_with_atts(cell, emp_code: str | None, emp_name: str | None, atts: list[DocumentAttachment] | None) -> bool:
        if not atts:
            return False
        valid_atts = [a for a in atts if a.file_name]
        if not valid_atts:
            return False

        cell.font = HYPERLINK_FONT
        if len(valid_atts) == 1:
            cell.hyperlink = build_attachment_zip_path(emp_code, emp_name, valid_atts[0].file_name)
            if not cell.value or str(cell.value).strip() == "":
                cell.value = "Xem file"
        else:
            cell.hyperlink = build_attachment_zip_folder(emp_code, emp_name)
            if not cell.value or str(cell.value).strip() == "":
                cell.value = f"Xem thư mục ({len(valid_atts)} file)"
        return True

    def _set_hyperlink_if_exists(cell, emp_code: str | None, emp_name: str | None, entity_type: str, entity_id: int | None) -> bool:
        if not entity_id:
            return False
        atts = attachment_map.get((entity_type, entity_id))
        return _set_hyperlink_with_atts(cell, emp_code, emp_name, atts)

    wb = Workbook()
    
    # Sheet 1: Danh sách Nhân sự
    ws1 = wb.active
    ws1.title = "Danh sách Nhân sự"
    headers1 = [
        "STT", "Mã NV", "Họ tên Latin", "Họ tên Trung Quốc", "Giới tính", "Ngày sinh", "Quốc tịch",
        "Số điện thoại", "Vị trí/Bộ phận", "Chức danh", "Số Hộ chiếu", "Hạn Hộ chiếu",
        "Ngày Đến VN", "Ngày Dự kiến sang", "Ngày Dự kiến về", "Ngày Thực tế về", "Hiện diện VN",
        "Chỗ ở hiện tại", "Loại hình", "Đăng ký ăn", "Loại Visa", "Hạn Visa", "Số ngày còn Visa",
        "Hạn Tạm trú", "Số ngày còn Tạm trú", "Số GPLĐ", "Cấp GPLĐ", "Hạn GPLĐ", "Số ngày còn GPLĐ",
        "Số HĐLĐ", "Loại HĐLĐ", "Hạn HĐLĐ", "Số ngày còn HĐLĐ", "Ghi chú",
    ]
    _apply_header_style(ws1, headers1)

    for idx, emp_eval in enumerate(evaluated, start=1):

        emp_obj = next((e for e in employees_db if e.id == emp_eval.id), None)

        # Calculate days remaining
        def days_rem(dt: datetime.date | None) -> int | None:
            if dt is None:
                return None
            return (dt - today).days

        visa_rem = days_rem(emp_eval.latest_visa_expiry)
        tamtru_rem = days_rem(emp_eval.latest_tamtru_expiry)
        gpld_rem = days_rem(emp_eval.latest_gpld_expiry)
        contract_rem = days_rem(emp_eval.latest_contract_expiry)

        # Latest Visa & TamTru
        latest_visa = None
        latest_tamtru = None
        if emp_obj and emp_obj.stays:
            all_visas = [v for s in emp_obj.stays for v in s.visas]
            if all_visas:
                sorted_visas = sorted(all_visas, key=lambda v: v.expiry_date or datetime.date.min, reverse=True)
                latest_visa = sorted_visas[0]

            all_tam_trus = [tt for s in emp_obj.stays for tt in s.tam_trus]
            if all_tam_trus:
                sorted_tts = sorted(all_tam_trus, key=lambda tt: tt.expiry_date or datetime.date.min, reverse=True)
                latest_tamtru = sorted_tts[0]

        # Latest WorkPermit & Contract details
        latest_wp = None
        if emp_obj and emp_obj.work_permits:
            sorted_wps = sorted(emp_obj.work_permits, key=lambda w: w.valid_to or datetime.date.min, reverse=True)
            latest_wp = sorted_wps[0]

        latest_contract = None
        if emp_obj and emp_obj.contracts:
            sorted_contracts = sorted(emp_obj.contracts, key=lambda c: c.end_date or datetime.date.min, reverse=True)
            latest_contract = sorted_contracts[0]


        row_data = [
            idx,
            emp_eval.employee_code or "",
            emp_eval.name_latin or "",
            emp_eval.name_chinese or "",
            emp_eval.gender or "",
            emp_eval.date_of_birth.strftime("%d/%m/%Y") if emp_eval.date_of_birth else "",
            emp_eval.nationality or "",
            emp_eval.phone or "",
            emp_eval.department or "",
            emp_eval.role or "",
            emp_eval.passport_number or "",
            emp_eval.passport_expiry.strftime("%d/%m/%Y") if emp_eval.passport_expiry else "",
            emp_eval.entry_date.strftime("%d/%m/%Y") if emp_eval.entry_date else "",
            emp_eval.expected_entry_date.strftime("%d/%m/%Y") if emp_eval.expected_entry_date else "",
            emp_eval.expected_exit_date.strftime("%d/%m/%Y") if emp_eval.expected_exit_date else "",
            emp_eval.actual_exit_date.strftime("%d/%m/%Y") if emp_eval.actual_exit_date else "",
            "Đang ở VN" if emp_eval.is_in_vietnam else "Đã về nước",
            emp_eval.current_room_number or ("Khách sạn" if emp_eval.is_in_vietnam else "–"),
            format_work_type(emp_eval.work_type),
            "Có" if emp_eval.is_in_vietnam else "Không",
            emp_eval.latest_visa_type or "",
            emp_eval.latest_visa_expiry.strftime("%d/%m/%Y") if emp_eval.latest_visa_expiry else "",
            visa_rem if visa_rem is not None else "",
            emp_eval.latest_tamtru_expiry.strftime("%d/%m/%Y") if emp_eval.latest_tamtru_expiry else "",
            tamtru_rem if tamtru_rem is not None else "",
            latest_wp.permit_number if (latest_wp and latest_wp.permit_number) else "",
            latest_wp.issue_type if (latest_wp and latest_wp.issue_type) else "",
            emp_eval.latest_gpld_expiry.strftime("%d/%m/%Y") if emp_eval.latest_gpld_expiry else "",
            gpld_rem if gpld_rem is not None else "",
            latest_contract.contract_number if (latest_contract and latest_contract.contract_number) else "",
            latest_contract.contract_type if (latest_contract and latest_contract.contract_type) else "",
            emp_eval.latest_contract_expiry.strftime("%d/%m/%Y") if emp_eval.latest_contract_expiry else "",
            contract_rem if contract_rem is not None else "",
            emp_eval.notes or "",
        ]
        ws1.append(row_data)

        # Highlight cell expiry dates & add hyperlinks if attachments exist
        curr_row = idx + 1

        def _get_atts_for_ids(entity_type: str, entity_ids: list[int]) -> list[DocumentAttachment]:
            res: list[DocumentAttachment] = []
            for eid in entity_ids:
                a_list = attachment_map.get((entity_type, eid))
                if a_list:
                    res.extend(a_list)
            return res

        # Passport hyperlink
        if emp_obj:
            passport_atts = attachment_map.get(("PASSPORT", emp_obj.id))
            _set_hyperlink_with_atts(ws1.cell(row=curr_row, column=11), emp_eval.employee_code, emp_eval.name_latin, passport_atts)

        # Visa hyperlink
        if emp_obj and emp_obj.stays:
            visa_ids = [v.id for s in emp_obj.stays for v in s.visas]
            visa_atts = _get_atts_for_ids("VISA", visa_ids)
            _set_hyperlink_with_atts(ws1.cell(row=curr_row, column=21), emp_eval.employee_code, emp_eval.name_latin, visa_atts)

        # TamTru hyperlink
        if emp_obj and emp_obj.stays:
            tt_ids = [tt.id for s in emp_obj.stays for tt in s.tam_trus]
            tt_atts = _get_atts_for_ids("TAM_TRU", tt_ids)
            _set_hyperlink_with_atts(ws1.cell(row=curr_row, column=24), emp_eval.employee_code, emp_eval.name_latin, tt_atts)

        # WorkPermit hyperlink
        if emp_obj and emp_obj.work_permits:
            wp_ids = [wp.id for wp in emp_obj.work_permits]
            wp_atts = _get_atts_for_ids("WORK_PERMIT", wp_ids)
            _set_hyperlink_with_atts(ws1.cell(row=curr_row, column=26), emp_eval.employee_code, emp_eval.name_latin, wp_atts)

        # Contract hyperlink
        if emp_obj and emp_obj.contracts:
            ct_ids = [c.id for c in emp_obj.contracts]
            ct_atts = _get_atts_for_ids("CONTRACT", ct_ids)
            _set_hyperlink_with_atts(ws1.cell(row=curr_row, column=30), emp_eval.employee_code, emp_eval.name_latin, ct_atts)

        # Helper to style warning
        def style_rem(col_idx: int, days: int | None):
            if days is None:
                return
            cell = ws1.cell(row=curr_row, column=col_idx)
            if days < 0:
                cell.fill = RED_FILL
                cell.font = RED_FONT
            elif days <= 30:
                cell.fill = YELLOW_FILL
                cell.font = YELLOW_FONT

        style_rem(23, visa_rem)
        style_rem(25, tamtru_rem)
        style_rem(29, gpld_rem)
        style_rem(33, contract_rem)

    _auto_fit_columns(ws1)

    # Sheet 2: Lịch sử Giấy tờ
    ws2 = wb.create_sheet(title="Lịch sử Giấy tờ")
    headers2 = ["STT", "Mã NV", "Họ tên Latin", "Loại Giấy tờ", "Số / Phân loại", "Ngày Cấp / Đăng ký", "Ngày Hết hạn", "Ghi chú"]
    _apply_header_style(ws2, headers2)

    s2_idx = 1
    for emp in employees_db:
        # Visas & TamTrus via stays
        for stay in emp.stays:
            for v in stay.visas:
                s2_row = s2_idx + 1
                ws2.append([
                    s2_idx, emp.employee_code or "", emp.name_latin, "VISA",
                    v.visa_type or "", v.entry_date.strftime("%d/%m/%Y") if v.entry_date else "",
                    v.expiry_date.strftime("%d/%m/%Y") if v.expiry_date else "", v.notes or ""
                ])
                _set_hyperlink_if_exists(ws2.cell(row=s2_row, column=5), emp.employee_code, emp.name_latin, "VISA", v.id)
                s2_idx += 1

            for tt in stay.tam_trus:
                s2_row = s2_idx + 1
                ws2.append([
                    s2_idx, emp.employee_code or "", emp.name_latin, "TAM_TRU",
                    "Đăng ký Tạm trú", tt.registration_date.strftime("%d/%m/%Y") if tt.registration_date else "",
                    tt.expiry_date.strftime("%d/%m/%Y") if tt.expiry_date else "", tt.notes or ""
                ])
                _set_hyperlink_if_exists(ws2.cell(row=s2_row, column=5), emp.employee_code, emp.name_latin, "TAM_TRU", tt.id)
                s2_idx += 1

        # WorkPermits
        for wp in emp.work_permits:
            s2_row = s2_idx + 1
            ws2.append([
                s2_idx, emp.employee_code or "", emp.name_latin, "GPLD",
                wp.permit_number or wp.issue_type or "", wp.valid_from.strftime("%d/%m/%Y") if wp.valid_from else "",
                wp.valid_to.strftime("%d/%m/%Y") if wp.valid_to else "", wp.notes or ""
            ])
            _set_hyperlink_if_exists(ws2.cell(row=s2_row, column=5), emp.employee_code, emp.name_latin, "WORK_PERMIT", wp.id)
            s2_idx += 1

        # Contracts
        for c in emp.contracts:
            s2_row = s2_idx + 1
            ws2.append([
                s2_idx, emp.employee_code or "", emp.name_latin, "CONTRACT",
                c.contract_number or c.contract_type or "", c.start_date.strftime("%d/%m/%Y") if c.start_date else "",
                c.end_date.strftime("%d/%m/%Y") if c.end_date else "", c.notes or ""
            ])
            _set_hyperlink_if_exists(ws2.cell(row=s2_row, column=5), emp.employee_code, emp.name_latin, "CONTRACT", c.id)
            s2_idx += 1


    _auto_fit_columns(ws2)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def generate_presence_accommodation_excel(db: Session, today: datetime.date | None = None) -> io.BytesIO:
    if today is None:
        today = datetime.date.today()

    employees_db = db.query(ForeignEmployee).order_by(ForeignEmployee.name_latin).all()
    evaluated = evaluate_employee_statuses(db, employees_db, today=today)

    wb = Workbook()

    # Sheet 1: Sơ đồ KTX & Khách sạn
    ws1 = wb.active
    ws1.title = "Sơ đồ KTX & Khách sạn"
    headers1 = [
        "STT", "Loại chỗ ở", "Tên Cơ sở / Phòng", "Số phòng KS", "Vị trí giường",
        "Mã NV", "Họ tên Latin", "Họ tên Trung Quốc", "Số Hộ chiếu", "Bộ phận",
        "Loại hình", "Đăng ký ăn", "Ngày vào ở", "Ngày dự kiến về"
    ]
    _apply_header_style(ws1, headers1)

    active_assigned = [e for e in evaluated if e.is_in_vietnam and e.current_room_number]
    for idx, emp_eval in enumerate(active_assigned, start=1):
        emp_obj = next((e for e in employees_db if e.id == emp_eval.id), None)
        active_stay = None
        if emp_obj:
            active_stay = next((s for s in emp_obj.stays if s.end_date is None or s.end_date >= today), None)

        acc_type = active_stay.accommodation_type if active_stay else "KTX"
        unit_name = emp_eval.current_room_number or ""
        bed_loc = active_stay.bed_location if active_stay else ""
        hotel_room = active_stay.hotel_room_number if active_stay else ""
        stay_type = active_stay.stay_type if active_stay else "CO_DINH"
        has_meals = "Có" if (active_stay and active_stay.has_meals) else "Không"
        start_date = active_stay.start_date.strftime("%d/%m/%Y") if (active_stay and active_stay.start_date) else ""
        exp_end = emp_eval.expected_exit_date.strftime("%d/%m/%Y") if emp_eval.expected_exit_date else ""

        ws1.append([
            idx, acc_type, unit_name, hotel_room, bed_loc,
            emp_eval.employee_code or "", emp_eval.name_latin, emp_eval.name_chinese or "",
            emp_eval.passport_number or "", emp_eval.department or "",
            format_work_type(stay_type), has_meals, start_date, exp_end
        ])
    _auto_fit_columns(ws1)

    # Sheet 2: Chưa xếp chỗ ở
    ws2 = wb.create_sheet(title="Chưa xếp chỗ ở")
    headers2 = ["STT", "Mã NV", "Họ tên Latin", "Họ tên Trung Quốc", "Số Hộ chiếu", "Bộ phận", "Ngày đến VN", "Ngày dự kiến về"]
    _apply_header_style(ws2, headers2)

    unassigned = [e for e in evaluated if e.is_in_vietnam and not e.current_room_number]
    for idx, emp_eval in enumerate(unassigned, start=1):
        ws2.append([
            idx, emp_eval.employee_code or "", emp_eval.name_latin, emp_eval.name_chinese or "",
            emp_eval.passport_number or "", emp_eval.department or "",
            emp_eval.entry_date.strftime("%d/%m/%Y") if emp_eval.entry_date else "",
            emp_eval.expected_exit_date.strftime("%d/%m/%Y") if emp_eval.expected_exit_date else ""
        ])
    _auto_fit_columns(ws2)

    # Sheet 3: Đã về nước
    ws3 = wb.create_sheet(title="Đã về nước")
    headers3 = ["STT", "Mã NV", "Họ tên Latin", "Họ tên Trung Quốc", "Số Hộ chiếu", "Bộ phận", "Ngày thực tế đã về", "Ngày dự kiến sang đợt tới"]
    _apply_header_style(ws3, headers3)

    exited = [e for e in evaluated if not e.is_in_vietnam]
    for idx, emp_eval in enumerate(exited, start=1):
        ws3.append([
            idx, emp_eval.employee_code or "", emp_eval.name_latin, emp_eval.name_chinese or "",
            emp_eval.passport_number or "", emp_eval.department or "",
            emp_eval.actual_exit_date.strftime("%d/%m/%Y") if emp_eval.actual_exit_date else "",
            emp_eval.expected_entry_date.strftime("%d/%m/%Y") if emp_eval.expected_entry_date else ""
        ])
    _auto_fit_columns(ws3)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def create_report_zip_package(
    excel_bytes: io.BytesIO,
    excel_filename: str,
    attachments: list[dict] | None = None
) -> io.BytesIO:
    zip_buf = io.BytesIO()
    with zipfile.ZipFile(zip_buf, "w", zipfile.ZIP_DEFLATED) as zf:
        # Write primary excel file
        zf.writestr(excel_filename, excel_bytes.getvalue())

        # Write document attachments if present
        if attachments:
            for item in attachments:
                code = item.get("employee_code")
                name = item.get("employee_name")
                file_name = item.get("file_name", "document.dat")
                content = item.get("content", b"")
                
                arc_path = build_attachment_zip_path(code, name, file_name)
                zf.writestr(arc_path, content)

    zip_buf.seek(0)
    return zip_buf
