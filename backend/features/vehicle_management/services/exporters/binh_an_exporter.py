from __future__ import annotations

from datetime import datetime
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Border, PatternFill, Side
from sqlalchemy.orm import Session

from features.vehicle_management.models import OwnershipGroup, VehicleDispatch
from features.vehicle_management.services.exporters.binh_an_sheet_builder import (
    build_binh_an_sheet1,
    build_binh_an_sheet2,
    is_border_trip,
)


def build_binh_an_excel_report(db: Session, from_date: str, to_date: str) -> BytesIO:
    """Generate Binh An 2-sheet Excel report using openpyxl."""
    wb = Workbook()
    thin_border = Border(
        left=Side(style="thin", color="000000"),
        right=Side(style="thin", color="000000"),
        top=Side(style="thin", color="000000"),
        bottom=Side(style="thin", color="000000"),
    )
    yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")

    try:
        f_dt = datetime.strptime(from_date, "%Y-%m-%d")
        t_dt = datetime.strptime(to_date, "%Y-%m-%d")
    except ValueError:
        f_dt = datetime.now()
        t_dt = datetime.now()

    month_str = f"{f_dt.month:02d}"
    year_str = f"{f_dt.year}"
    date_range_str = f"{f_dt.strftime('%d/%m/%Y')} đến {t_dt.strftime('%d/%m/%Y')}"

    dispatches = (
        db.query(VehicleDispatch)
        .filter(
            VehicleDispatch.ownership_group == OwnershipGroup.OUTSOURCED,
            VehicleDispatch.dispatch_date >= from_date,
            VehicleDispatch.dispatch_date <= to_date,
        )
        .order_by(VehicleDispatch.dispatch_date.asc(), VehicleDispatch.id.asc())
        .all()
    )

    border_trips = [d for d in dispatches if is_border_trip(d)]
    normal_trips = [d for d in dispatches if not is_border_trip(d)]

    ws1 = wb.active
    ws1.title = "BINH AN 08"
    build_binh_an_sheet1(ws1, month_str, year_str, date_range_str, normal_trips, border_trips, thin_border, yellow_fill)

    ws2 = wb.create_sheet(title="HỮU NGHỊ BÌNH AN")
    build_binh_an_sheet2(ws2, month_str, year_str, date_range_str, border_trips, thin_border, yellow_fill)

    out = BytesIO()
    wb.save(out)
    out.seek(0)
    return out
