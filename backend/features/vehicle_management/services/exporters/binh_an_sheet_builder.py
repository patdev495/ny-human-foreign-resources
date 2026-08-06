from __future__ import annotations

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet
from features.vehicle_management.models import VehicleDispatch


def get_bilingual_purpose(text: str) -> str:
    t = (text or "").lower()
    if "làm đêm" in t or "đêm" in t:
        return "Đi làm đêm\n上夜班"
    if "tan làm" in t or "về" in t:
        return "Tan làm \n下班"
    if "đi làm" in t or ("đi" in t and "làm" in t):
        return "Đi làm\n上班"
    if "đi ăn" in t or "ăn" in t:
        return "đi ăn\n吃饭"
    if "sân bay" in t or "nội bài" in t:
        return "san bay\n机场"
    if "cửa khẩu" in t or "hữu nghị" in t or "友谊关" in t:
        return "cửa khẩu\n友谊关"
    if "đình trám" in t:
        return "đình trám\n工业区"
    if "bắc giang" in t:
        return "bắc giang\n北江"
    if "giao hàng" in t:
        return "giao hàng\n运货"
    return text if text else "Đi làm\n上班"


def get_bilingual_route(pickup: str, dropoff: str) -> str:
    p = (pickup or "").lower()
    d = (dropoff or "").lower()

    if ("kí túc" in p or "ktx" in p) and ("công ty" in d or "nienyi" in d or "cn09" in d or "cn15" in d):
        return "kí túc -công ty \n宿舍-公司"
    if ("công ty" in p or "nienyi" in p or "cn09" in p or "cn15" in p) and ("kí túc" in d or "ktx" in d):
        return "công ty - kí túc\n公司-宿舍"
    if ("kí túc" in p or "ktx" in p) and ("cửa khẩu" in d or "hữu nghị" in d or "友谊关" in d):
        return "kí túc -cửa khẩu \n宿舍- 友谊关"
    if ("cửa khẩu" in p or "hữu nghị" in p or "友谊关" in p) and ("kí túc" in d or "ktx" in d):
        return "cửa khẩu_ ky túc\n友谊关_宿舍"
    if "sân bay" in p or "sân bay" in d or "nội bài" in p or "nội bài" in d:
        return "sân bay - kí túc\n机场-宿舍"

    return f"{pickup or ''} - {dropoff or ''}"


def is_border_trip(d: VehicleDispatch) -> bool:
    text = f"{d.pickup_location or ''} {d.dropoff_location or ''} {d.notes or ''} {d.vehicle_name or ''}".lower()
    return any(k in text for k in ["cửa khẩu", "hữu nghị", "友谊关", "hnq"])


def build_binh_an_sheet1(
    ws1: Worksheet,
    month_str: str,
    year_str: str,
    date_range_str: str,
    normal_trips: list[VehicleDispatch],
    border_trips: list[VehicleDispatch],
    thin_border: Border,
    yellow_fill: PatternFill,
) -> None:
    ws1.views.sheetView[0].showGridLines = True
    col_widths_1 = {"A": 11.57, "B": 20.57, "C": 19.86, "D": 43.14, "E": 15.71, "F": 14.71, "G": 22.29, "H": 18.86, "I": 62.57}
    for col, width in col_widths_1.items():
        ws1.column_dimensions[col].width = width

    ws1.row_dimensions[1].height = 19.5
    ws1.row_dimensions[2].height = 19.5
    ws1.row_dimensions[3].height = 19.5
    ws1.row_dimensions[4].height = 39.75
    ws1.row_dimensions[5].height = 43.5
    ws1.row_dimensions[6].height = 33.0
    ws1.row_dimensions[7].height = 27.0
    ws1.row_dimensions[8].height = 68.25

    ws1.merge_cells("A1:E1")
    ws1["A1"] = "CÔNG TY CP THƯƠNG MẠI ĐẦU TƯ VÀ PHÁT TRIỂN DỊCH VỤ DU LỊCH VẬN TẢI BÌNH AN"
    ws1["A1"].font = Font(name="Calibri Light", size=15, bold=True)
    ws1["A1"].alignment = Alignment(horizontal="left", vertical="center")

    ws1.merge_cells("A2:E2")
    ws1["A2"] = "ĐC: Ngõ Đền, TDP Hoàng Mai 3, Phường Nếnh, tỉnh Bắc Ninh, Việt Nam  "
    ws1["A2"].font = Font(name="Calibri Light", size=15, bold=True)
    ws1["A2"].alignment = Alignment(horizontal="left", vertical="center")

    ws1.merge_cells("A3:E3")
    ws1["A3"] = "Mã số Thuế: 2400960417"
    ws1["A3"].font = Font(name="Calibri Light", size=15, bold=True)
    ws1["A3"].alignment = Alignment(horizontal="left", vertical="center")

    ws1.merge_cells("A4:I4")
    ws1["A4"] = "BẢNG ĐỐI CHIẾU KHỐI LƯỢNG VÀ GIÁ TRỊ SỬ DỤNG XE Ô TÔ \n租车使用对账明细表"
    ws1["A4"].font = Font(name="Calibri Light", size=15, bold=True)
    ws1["A4"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws1.merge_cells("A5:I5")
    ws1["A5"] = f"Tháng {month_str} năm {year_str}\n{year_str} 年 {month_str}月"
    ws1["A5"].font = Font(name="Calibri Light", size=15, bold=True)
    ws1["A5"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws1.merge_cells("A6:I6")
    ws1["A6"] = f"Bảng kê từ ngày {date_range_str} kèm theo hoá đơn số:   kí hiệu:  ngày:"
    ws1["A6"].font = Font(name="Calibri Light", size=15, bold=True)
    ws1["A6"].alignment = Alignment(horizontal="center", vertical="center")

    ws1.merge_cells("A7:I7")
    ws1["A7"] = "   Khách hàng客户: CÔNG TY TNHH CÔNG NGHIỆP NIENYI VIỆT NAM                                                                                                                       "
    ws1["A7"].font = Font(name="Calibri Light", size=15, bold=True)
    ws1["A7"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    headers_1 = ["序号\nStt", "日期\nNgày", "用途\nMục đích", "行程\nHành trình", "车\nxe", "等费\nPHÍ ĐỢI\n(VNĐ)", "费车(VND)\nSố Tiền", "车牌 Biển số xe", "备注\nGhi Chú"]
    for c_idx, h_text in enumerate(headers_1, start=1):
        cell = ws1.cell(row=8, column=c_idx, value=h_text)
        cell.font = Font(name="Calibri Light", size=11 if c_idx == 6 else 15, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.fill = yellow_fill
        cell.border = thin_border

    row_idx = 9
    stt = 1
    subtotal_rows = []
    char_codes = ["A", "B", "C", "D", "E", "F", "G", "H"]
    char_idx = 0

    grouped_trips: dict[str, list[VehicleDispatch]] = {}
    for d in normal_trips:
        plate = d.license_plate or d.vehicle_name or "Khác"
        grouped_trips.setdefault(plate, []).append(d)

    for plate, trips in grouped_trips.items():
        block_code = char_codes[char_idx % len(char_codes)]
        char_idx += 1
        start_row = row_idx

        for d in trips:
            ws1.row_dimensions[row_idx].height = 42.0
            purpose_bilingual = get_bilingual_purpose(d.notes or "")
            route_bilingual = get_bilingual_route(d.pickup_location or "", d.dropoff_location or "")

            ws1.cell(row=row_idx, column=1, value=stt).alignment = Alignment(horizontal="center", vertical="center")
            ws1.cell(row=row_idx, column=2, value=d.dispatch_date).alignment = Alignment(horizontal="center", vertical="center")
            ws1.cell(row=row_idx, column=3, value=purpose_bilingual).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws1.cell(row=row_idx, column=4, value=route_bilingual).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws1.cell(row=row_idx, column=5, value=d.vehicle_name or "").alignment = Alignment(horizontal="center", vertical="center")

            wait_fee = d.waiting_hours * 30000 if d.waiting_hours else ""
            cell_f = ws1.cell(row=row_idx, column=6, value=wait_fee)
            cell_f.alignment = Alignment(horizontal="right", vertical="center")
            if wait_fee:
                cell_f.number_format = "#,##0"

            cell_g = ws1.cell(row=row_idx, column=7, value=d.cost)
            cell_g.alignment = Alignment(horizontal="right", vertical="center")
            cell_g.number_format = "#,##0"

            ws1.cell(row=row_idx, column=8, value=d.license_plate or d.vehicle_name).alignment = Alignment(horizontal="center", vertical="center")

            note_str = f"{d.passenger_count} 个人=>"
            if d.pickup_time:
                note_str += f"{d.pickup_time}"
            if d.passenger_name:
                note_str += f" {d.passenger_name}"
            ws1.cell(row=row_idx, column=9, value=note_str).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

            for c in range(1, 10):
                cell = ws1.cell(row=row_idx, column=c)
                cell.font = Font(name="Calibri Light", size=14, bold=True)
                cell.border = thin_border

            stt += 1
            row_idx += 1

        sub_r = row_idx
        ws1.row_dimensions[sub_r].height = 42.0
        ws1.cell(row=sub_r, column=1, value=block_code).alignment = Alignment(horizontal="center", vertical="center")
        ws1.cell(row=sub_r, column=2, value=f"TỔNG XE {plate}").alignment = Alignment(horizontal="left", vertical="center")
        cell_sum = ws1.cell(row=sub_r, column=7, value=f"=SUM(G{start_row}:G{sub_r - 1})")
        cell_sum.alignment = Alignment(horizontal="right", vertical="center")
        cell_sum.number_format = "#,##0"

        for c in range(1, 10):
            cell = ws1.cell(row=sub_r, column=c)
            cell.font = Font(name="Calibri Light", size=14, bold=True)
            cell.border = thin_border

        subtotal_rows.append(sub_r)
        row_idx += 1

    if border_trips:
        block_code = char_codes[char_idx % len(char_codes)]
        char_idx += 1
        b_total = sum(d.cost for d in border_trips)
        sub_r = row_idx
        ws1.row_dimensions[sub_r].height = 42.0
        ws1.cell(row=sub_r, column=1, value=block_code).alignment = Alignment(horizontal="center", vertical="center")
        ws1.cell(row=sub_r, column=2, value="TỔNG XE CỬA KHẨU").alignment = Alignment(horizontal="left", vertical="center")
        cell_sum = ws1.cell(row=sub_r, column=7, value=b_total)
        cell_sum.alignment = Alignment(horizontal="right", vertical="center")
        cell_sum.number_format = "#,##0"

        for c in range(1, 10):
            cell = ws1.cell(row=sub_r, column=c)
            cell.font = Font(name="Calibri Light", size=14, bold=True)
            cell.border = thin_border

        subtotal_rows.append(sub_r)
        row_idx += 1

    grand_r = row_idx
    ws1.row_dimensions[grand_r].height = 42.0
    ws1.cell(row=grand_r, column=1, value="TỔNG PHẢI TRẢ (VNĐ)").alignment = Alignment(horizontal="left", vertical="center")
    if subtotal_rows:
        sum_formula = "=" + "+".join([f"G{r}" for r in subtotal_rows])
        cell_grand = ws1.cell(row=grand_r, column=7, value=sum_formula)
        cell_grand.alignment = Alignment(horizontal="right", vertical="center")
        cell_grand.number_format = "#,##0"

    for c in range(1, 10):
        cell = ws1.cell(row=grand_r, column=c)
        cell.font = Font(name="Calibri Light", size=14, bold=True)
        cell.border = thin_border


def build_binh_an_sheet2(
    ws2: Worksheet,
    month_str: str,
    year_str: str,
    date_range_str: str,
    border_trips: list[VehicleDispatch],
    thin_border: Border,
    yellow_fill: PatternFill,
) -> None:
    ws2.views.sheetView[0].showGridLines = True
    col_widths_2 = {"A": 9.86, "B": 18.29, "C": 19.57, "D": 32.57, "E": 18.29, "F": 23.29, "G": 22.86, "H": 26.57, "I": 54.71}
    for col, width in col_widths_2.items():
        ws2.column_dimensions[col].width = width

    for r_idx, h in enumerate([21.0, 21.0, 21.0, 26.25, 34.5, 36.0, 21.0, 37.5], start=1):
        ws2.row_dimensions[r_idx].height = h

    ws2.merge_cells("A1:E1")
    ws2["A1"] = "CÔNG TY CP THƯƠNG MẠI ĐẦU TƯ VÀ PHÁT TRIỂN DỊCH VỤ DU LỊCH VẬN TẢI BÌNH AN"
    ws2["A1"].font = Font(name="Calibri Light", size=16, bold=False)
    ws2["A1"].alignment = Alignment(horizontal="left", vertical="center")

    ws2.merge_cells("A2:E2")
    ws2["A2"] = "ĐC: Ngõ Đền, TDP Hoàng Mai 3, Phường Nếnh, tỉnh Bắc Ninh, Việt Nam  "
    ws2["A2"].font = Font(name="Calibri Light", size=16, bold=False)
    ws2["A2"].alignment = Alignment(horizontal="left", vertical="center")

    ws2.merge_cells("A3:E3")
    ws2["A3"] = "Mã số Thuế: 2400960417"
    ws2["A3"].font = Font(name="Calibri Light", size=16, bold=False)
    ws2["A3"].alignment = Alignment(horizontal="left", vertical="center")

    ws2.merge_cells("A4:H4")
    ws2["A4"] = "BẢNG ĐỐI CHIẾU KHỐI LƯỢNG VÀ GIÁ TRỊ SỬ DỤNG XE Ô TÔ \n租车使用对账明细表"
    ws2["A4"].font = Font(name="Calibri Light", size=20, bold=False)
    ws2["A4"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws2.merge_cells("A5:H5")
    ws2["A5"] = f"Tháng {month_str} năm {year_str}\n{year_str} 年 {month_str}月"
    ws2["A5"].font = Font(name="Calibri Light", size=16, bold=False)
    ws2["A5"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws2.merge_cells("A6:H6")
    ws2["A6"] = f"Bảng kê từ ngày {date_range_str}"
    ws2["A6"].font = Font(name="Calibri Light", size=16, bold=False)
    ws2["A6"].alignment = Alignment(horizontal="center", vertical="center")

    ws2.merge_cells("A7:H7")
    ws2["A7"] = "   Khách hàng客户: CÔNG TY TNHH CÔNG NGHIỆP NIENYI VIỆT NAM                                                                                                                       "
    ws2["A7"].font = Font(name="Calibri Light", size=16, bold=False)
    ws2["A7"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    headers_2 = ["序号\nStt", "日期\nNgày", "用途\nMục đích", "行程\nHành trình", "车\nxe", "费车(VND)\nSố Tiền", "车牌 Biển số xe", "备注\nGhi Chú", "HÀNH KHÁCH"]
    for c_idx, h_text in enumerate(headers_2, start=1):
        cell = ws2.cell(row=8, column=c_idx, value=h_text)
        cell.font = Font(name="Calibri Light", size=14, bold=False)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.fill = yellow_fill
        cell.border = thin_border

    r2_idx = 9
    stt2 = 1
    for d in border_trips:
        ws2.row_dimensions[r2_idx].height = 47.25
        purpose_bilingual = get_bilingual_purpose(d.notes or "cửa khẩu")
        route_bilingual = get_bilingual_route(d.pickup_location or "", d.dropoff_location or "")

        ws2.cell(row=r2_idx, column=1, value=stt2).alignment = Alignment(horizontal="center", vertical="center")
        ws2.cell(row=r2_idx, column=2, value=d.dispatch_date).alignment = Alignment(horizontal="center", vertical="center")
        ws2.cell(row=r2_idx, column=3, value=purpose_bilingual).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws2.cell(row=r2_idx, column=4, value=route_bilingual).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws2.cell(row=r2_idx, column=5, value=d.vehicle_name or "4座chỗ").alignment = Alignment(horizontal="center", vertical="center")

        cell_f = ws2.cell(row=r2_idx, column=6, value=d.cost)
        cell_f.alignment = Alignment(horizontal="right", vertical="center")
        cell_f.number_format = "#,##0"

        ws2.cell(row=r2_idx, column=7, value=d.license_plate or d.vehicle_name).alignment = Alignment(horizontal="center", vertical="center")
        ws2.cell(row=r2_idx, column=8, value=f"{d.pickup_time or ''}_{d.return_time or ''}").alignment = Alignment(horizontal="center", vertical="center")
        ws2.cell(row=r2_idx, column=9, value=f"{d.passenger_count}个人 => {d.passenger_name or ''}").alignment = Alignment(horizontal="left", vertical="center")

        for c in range(1, 10):
            cell = ws2.cell(row=r2_idx, column=c)
            cell.font = Font(name="Calibri Light", size=14, bold=False)
            cell.border = thin_border

        stt2 += 1
        r2_idx += 1

    if border_trips:
        ws2.row_dimensions[r2_idx].height = 47.25
        ws2.cell(row=r2_idx, column=1, value="F").alignment = Alignment(horizontal="center", vertical="center")
        ws2.cell(row=r2_idx, column=2, value="TỔNG XE CỬA KHẨU").alignment = Alignment(horizontal="left", vertical="center")
        cell_sum = ws2.cell(row=r2_idx, column=6, value=f"=SUM(F9:F{r2_idx - 1})")
        cell_sum.alignment = Alignment(horizontal="right", vertical="center")
        cell_sum.number_format = "#,##0"

        for c in range(1, 10):
            cell = ws2.cell(row=r2_idx, column=c)
            cell.font = Font(name="Calibri Light", size=14, bold=True)
            cell.border = thin_border
