from __future__ import annotations

from datetime import datetime, timedelta
from io import BytesIO
from pathlib import Path
from typing import Optional
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from features.vehicle_management.models import MonthlyVehicleContract, VehicleDispatch


def get_report_dir() -> Optional[Path]:
    possible_dirs = [
        Path.cwd() / "report",
        Path.cwd().parent / "report",
        Path(__file__).resolve().parents[4] / "report",
        Path(__file__).resolve().parents[3] / "report",
        Path(__file__).resolve().parents[2] / "report",
        Path(__file__).resolve().parents[1] / "report",
    ]
    for d in possible_dirs:
        if d.exists():
            return d
    return None


def find_sample_template(license_plate: str) -> Optional[Path]:
    if not license_plate:
        return None

    base_dir = get_report_dir()
    if not base_dir:
        return None

    clean_plate = license_plate.replace("-", "").replace(".", "").upper()
    for f in base_dir.glob("*.xlsx"):
        clean_fname = f.name.replace("-", "").replace(".", "").upper()
        if ("81988" in clean_plate and "81988" in clean_fname) or \
           ("36900" in clean_plate and "36900" in clean_fname) or \
           ("10378" in clean_plate and "10378" in clean_fname):
            return f
    return None



def populate_vehicle_from_template(
    template_path: Path,
    db: Session,
    contract: MonthlyVehicleContract,
    from_date: str,
    to_date: str,
) -> BytesIO:
    v = contract.vehicle
    wb = load_workbook(template_path)

    sheet_name = "BẢNG KÊ CHI TIẾT XUẤT VAT 07"
    if sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
    else:
        ws = wb.active

    try:
        f_dt = datetime.strptime(from_date, "%Y-%m-%d")
        t_dt = datetime.strptime(to_date, "%Y-%m-%d")
    except ValueError:
        f_dt = datetime.now()
        t_dt = datetime.now()

    month_str = f"{f_dt.month:02d}"
    year_str = f"{f_dt.year}"

    # Update Title Header in template
    ws["A4"] = f"Tháng {month_str} năm {year_str}\n{year_str} 年{month_str}月"
    ws["A6"] = f"Bảng kê số {month_str}{year_str} ngày 25/{month_str}/{year_str} kèm hóa đơn số:            kí hiệu:"

    # Clear all sample data in rows 12..41 before writing dynamic range
    for r in range(12, 42):
        for c in range(1, 30):
            cell = ws.cell(row=r, column=c)
            val = str(cell.value or "")
            if not val.startswith("="):
                cell.value = None

    # Query dispatches for this vehicle in range
    dispatches = (
        db.query(VehicleDispatch)
        .filter(
            VehicleDispatch.vehicle_id == v.id,
            VehicleDispatch.dispatch_date >= from_date,
            VehicleDispatch.dispatch_date <= to_date,
        )
        .all()
    )

    daily_map = {}
    for d in dispatches:
        daily_map.setdefault(d.dispatch_date, []).append(d)

    # Fill daily rows starting at row 12
    curr_dt = f_dt
    day_idx = 0

    while curr_dt <= t_dt and (12 + day_idx) <= 41:
        r = 12 + day_idx
        d_str = curr_dt.strftime("%Y-%m-%d")

        ws.cell(row=r, column=1).value = d_str
        trips = daily_map.get(d_str, [])

        if trips:
            trips_sorted = sorted(trips, key=lambda t: (t.pickup_time or "00:00", t.id))
            f_trip = trips_sorted[0]
            l_trip = trips_sorted[-1]

            s_km = f_trip.odometer_km or f_trip.start_km
            e_km = l_trip.end_km or l_trip.odometer_km

            s_time = f_trip.pickup_time or "07:00"
            e_time = l_trip.return_time or l_trip.pickup_time or "19:00"

            if s_km is not None: ws.cell(row=r, column=2).value = s_km
            if e_km is not None: ws.cell(row=r, column=3).value = e_km
            ws.cell(row=r, column=4).value = f"=+C{r}-B{r}"

            ws.cell(row=r, column=11).value = s_time
            ws.cell(row=r, column=12).value = e_time

            day_toll = sum(getattr(t, "toll_fee", 0.0) or 0.0 for t in trips)
            day_meal_cnt = sum(getattr(t, "meal_count", 0) or 0 for t in trips)
            day_overnight_cnt = sum(getattr(t, "overnight_count", 0) or 0 for t in trips)

            if day_overnight_cnt > 0:
                ws.cell(row=r, column=18).value = day_overnight_cnt * (contract.overnight_fee or 300000)

            if day_meal_cnt > 0:
                ws.cell(row=r, column=25).value = day_meal_cnt
                ws.cell(row=r, column=26).value = f"=Y{r}*50000"

            if day_toll > 0:
                ws.cell(row=r, column=27).value = day_toll

        curr_dt += timedelta(days=1)
        day_idx += 1

        curr_dt += timedelta(days=1)

    # Update sheet ĐNTT if present
    if "ĐNTT" in wb.sheetnames:
        ws_dntt = wb["ĐNTT"]
        ws_dntt["A1"] = f"{year_str} 年{month_str}月份的对账单\nBẢNG ĐỐI CHIẾU CÔNG NỢ CƯỚC XE THÁNG {month_str}/{year_str}"
        if ws_dntt.max_row >= 29:
            ws_dntt.cell(row=29, column=2).value = f"Bắc Ninh, ngày 25 tháng {month_str} năm {year_str}"

    out = BytesIO()
    wb.save(out)
    out.seek(0)
    return out


def find_binh_an_template() -> Optional[Path]:
    base_dir = get_report_dir()
    if not base_dir:
        return None

    for name in ["template_binh_an.xlsx", "NIENYI-BÌNH AN 07.xlsx", "NIENYI-BINH AN 07.xlsx"]:
        p = base_dir / name
        if p.exists():
            return p
    return None



def get_bilingual_purpose(text: str) -> str:
    t = (text or '').lower()
    if 'làm đêm' in t or 'đêm' in t: return 'Đi làm đêm\n上夜班'
    if 'tan làm' in t or 'về' in t: return 'Tan làm \n下班'
    if 'đi làm' in t or ('đi' in t and 'làm' in t): return 'Đi làm\n上班'
    if 'đi ăn' in t or 'ăn' in t: return 'đi ăn\n吃饭'
    if 'sân bay' in t or 'nội bài' in t: return 'san bay\n机场'
    if 'cửa khẩu' in t or 'hữu nghị' in t or '友谊关' in t: return 'cửa khẩu\n友谊关'
    if 'đình trám' in t: return 'đình trám\n工业区'
    if 'bắc giang' in t: return 'bắc giang\n北江'
    if 'giao hàng' in t: return 'giao hàng\n运货'
    return text if text else 'Đi làm\n上班'


def get_bilingual_route(pickup: str, dropoff: str) -> str:
    p = (pickup or '').lower()
    d = (dropoff or '').lower()

    if ('kí túc' in p or 'ktx' in p) and ('công ty' in d or 'nienyi' in d or 'cn09' in d or 'cn15' in d):
        return 'kí túc -công ty \n宿舍-公司'
    if ('công ty' in p or 'nienyi' in p or 'cn09' in p or 'cn15' in p) and ('kí túc' in d or 'ktx' in d):
        return 'công ty - kí túc\n公司-宿舍'
    if ('kí túc' in p or 'ktx' in p) and ('cửa khẩu' in d or 'hữu nghị' in d or '友谊关' in d):
        return 'kí túc -cửa khẩu \n宿舍- 友谊关'
    if ('cửa khẩu' in p or 'hữu nghị' in p or '友谊关' in p) and ('kí túc' in d or 'ktx' in d):
        return 'cửa khẩu_ ky túc\n友谊关_宿舍'
    if 'sân bay' in p or 'sân bay' in d or 'nội bài' in p or 'nội bài' in d:
        return 'sân bay - kí túc\n机场-宿舍'

    p_str = pickup or ''
    d_str = dropoff or ''
    return f'{p_str} - {d_str}'


def populate_binh_an_from_template(
    template_path: Path,
    db: Session,
    from_date: str,
    to_date: str,
) -> BytesIO:
    from openpyxl.cell.cell import MergedCell
    from openpyxl.styles import Alignment, Font, PatternFill, Side, Border
    from features.vehicle_management.models import OwnershipGroup, VehicleDispatch

    wb = load_workbook(template_path)

    ws1 = wb["BINH AN 08"] if "BINH AN 08" in wb.sheetnames else wb.worksheets[0]
    ws2 = wb["HỮU NGHỊ BÌNH AN"] if "HỮU NGHỊ BÌNH AN" in wb.sheetnames else (wb.worksheets[1] if len(wb.worksheets) > 1 else None)

    def set_cell(ws, r, c, val):
        if ws is None:
            return
        cell = ws.cell(row=r, column=c)
        if not isinstance(cell, MergedCell):
            cell.value = val

    try:
        f_dt = datetime.strptime(from_date, "%Y-%m-%d")
        t_dt = datetime.strptime(to_date, "%Y-%m-%d")
    except ValueError:
        f_dt = datetime.now()
        t_dt = datetime.now()

    month_str = f"{f_dt.month:02d}"
    year_str = f"{f_dt.year}"

    set_cell(ws1, 5, 1, f"Tháng {month_str} năm {year_str}\n{year_str} 年 {month_str}月")
    set_cell(ws1, 6, 1, f"Bảng kê từ ngày {f_dt.strftime('%d/%m/%Y')} đến {t_dt.strftime('%d/%m/%Y')} kèm theo hoá đơn số:   kí hiệu:  ngày:")

    if ws2:
        set_cell(ws2, 5, 1, f"Tháng {month_str} năm {year_str}\n{year_str} 年 {month_str}月")
        set_cell(ws2, 6, 1, f"Bảng kê từ ngày {f_dt.strftime('%d/%m/%Y')} đến {t_dt.strftime('%d/%m/%Y')}")

    dispatches = (
        db.query(VehicleDispatch)
        .filter(
            VehicleDispatch.ownership_group == OwnershipGroup.OUTSOURCED,
            VehicleDispatch.dispatch_date >= from_date,
            VehicleDispatch.dispatch_date <= to_date,
        )
        .order_by(VehicleDispatch.dispatch_date.asc(), VehicleDispatch.id.asc())
        .all()
    )

    def is_border_trip(d: VehicleDispatch) -> bool:
        text = f"{d.pickup_location or ''} {d.dropoff_location or ''} {d.notes or ''} {d.vehicle_name or ''}".lower()
        return any(k in text for k in ["cửa khẩu", "hữu nghị", "友谊关", "hnq"])

    border_trips = [d for d in dispatches if is_border_trip(d)]
    normal_trips = [d for d in dispatches if not is_border_trip(d)]

    if dispatches:
        for r in range(9, ws1.max_row + 1):
            for c in range(1, 10):
                set_cell(ws1, r, c, None)

        row_idx = 9
        stt = 1
        subtotal_rows = []
        char_codes = ["A", "B", "C", "D", "E", "F", "G", "H"]
        char_idx = 0

        grouped_trips = {}
        for d in normal_trips:
            plate = d.license_plate or d.vehicle_name or "Khác"
            grouped_trips.setdefault(plate, []).append(d)

        for plate, trips in grouped_trips.items():
            block_code = char_codes[char_idx % len(char_codes)]
            char_idx += 1
            start_row = row_idx

            for d in trips:
                purpose_bilingual = get_bilingual_purpose(d.notes or "")
                route_bilingual = get_bilingual_route(d.pickup_location or "", d.dropoff_location or "")

                set_cell(ws1, row_idx, 1, stt)
                set_cell(ws1, row_idx, 2, d.dispatch_date)
                set_cell(ws1, row_idx, 3, purpose_bilingual)
                set_cell(ws1, row_idx, 4, route_bilingual)
                set_cell(ws1, row_idx, 5, d.vehicle_name or "")
                set_cell(ws1, row_idx, 6, d.waiting_hours * 30000 if d.waiting_hours else "")
                set_cell(ws1, row_idx, 7, d.cost)
                set_cell(ws1, row_idx, 8, d.license_plate or d.vehicle_name)

                note_str = f"{d.passenger_count} 个人=>"
                if d.pickup_time:
                    note_str += f"{d.pickup_time}"
                if d.passenger_name:
                    note_str += f" {d.passenger_name}"
                set_cell(ws1, row_idx, 9, note_str)

                stt += 1
                row_idx += 1

            sub_r = row_idx
            set_cell(ws1, sub_r, 1, block_code)
            set_cell(ws1, sub_r, 2, f"TỔNG XE {plate}")
            set_cell(ws1, sub_r, 7, f"=SUM(G{start_row}:G{sub_r - 1})")
            subtotal_rows.append(sub_r)
            row_idx += 1

        if border_trips:
            block_code = char_codes[char_idx % len(char_codes)]
            char_idx += 1
            b_total = sum(d.cost for d in border_trips)
            set_cell(ws1, row_idx, 1, block_code)
            set_cell(ws1, row_idx, 2, "TỔNG XE CỬA KHẨU")
            set_cell(ws1, row_idx, 7, b_total)
            subtotal_rows.append(row_idx)
            row_idx += 1

        grand_r = row_idx
        set_cell(ws1, grand_r, 1, "TỔNG PHẢI TRẢ (VNĐ)")
        if subtotal_rows:
            sum_formula = "=" + "+".join([f"G{r}" for r in subtotal_rows])
            set_cell(ws1, grand_r, 7, sum_formula)

        if ws2:
            for r in range(9, ws2.max_row + 1):
                for c in range(1, 10):
                    set_cell(ws2, r, c, None)

            r2_idx = 9
            stt2 = 1
            for d in border_trips:
                purpose_bilingual = get_bilingual_purpose(d.notes or "cửa khẩu")
                route_bilingual = get_bilingual_route(d.pickup_location or "", d.dropoff_location or "")

                set_cell(ws2, r2_idx, 1, stt2)
                set_cell(ws2, r2_idx, 2, d.dispatch_date)
                set_cell(ws2, r2_idx, 3, purpose_bilingual)
                set_cell(ws2, r2_idx, 4, route_bilingual)
                set_cell(ws2, r2_idx, 5, d.vehicle_name or "4座chỗ")
                set_cell(ws2, r2_idx, 6, d.cost)
                set_cell(ws2, r2_idx, 7, d.license_plate or d.vehicle_name)
                set_cell(ws2, r2_idx, 8, f"{d.pickup_time or ''}_{d.return_time or ''}")
                set_cell(ws2, r2_idx, 9, f"{d.passenger_count}个人 => {d.passenger_name or ''}")

                stt2 += 1
                r2_idx += 1

            if border_trips:
                set_cell(ws2, r2_idx, 1, "F")
                set_cell(ws2, r2_idx, 2, "TỔNG XE CỬA KHẨU")
                set_cell(ws2, r2_idx, 6, f"=SUM(F9:F{r2_idx - 1})")

    out = BytesIO()
    wb.save(out)
    out.seek(0)
    return out



