from __future__ import annotations

import io
import datetime
import pytest
from openpyxl import load_workbook
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from main import app
from core.database import Base, get_db
from features.hr_foreign.models import (
    ForeignEmployee,
    Room,
    Stay,
    Visa,
    TamTru,
    WorkPermit,
    Contract,
    DocumentAttachment,
)
from features.hr_foreign.excel_exporter import (
    generate_legal_profile_excel,
    generate_presence_accommodation_excel,
    create_report_zip_package,
    build_attachment_zip_path,
)



@pytest.fixture
def db_session():
    engine = create_engine(
        "sqlite:///:memory:",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(bind=engine)
    TestingSessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)
    session = TestingSessionLocal()
    try:
        yield session
    finally:
        session.close()


@pytest.fixture
def client(db_session):
    def override_get_db():
        yield db_session

    app.dependency_overrides[get_db] = override_get_db
    with TestClient(app) as c:
        yield c
    app.dependency_overrides.clear()


def test_generate_legal_profile_excel(db_session):
    # Setup test employee with stays, visas, tam trus, work permits, contracts
    emp = ForeignEmployee(
        employee_code="NY1001",
        name_latin="WANG LEI",
        name_chinese="王伟",
        gender="Nam",
        nationality="Trung Quoc",
        date_of_birth=datetime.date(1990, 1, 1),
        passport_number="E12345678",
        passport_expiry=datetime.date(2030, 1, 1),
        department="Ky Thuat",
        role="Ky su",
        entry_date=datetime.date(2026, 1, 10),
    )
    db_session.add(emp)
    db_session.flush()

    stay = Stay(
        employee_id=emp.id,
        accommodation_type="KTX",
        stay_type="CO_DINH",
        has_meals=True,
        start_date=datetime.date(2026, 1, 10),
    )
    db_session.add(stay)
    db_session.flush()

    visa = Visa(
        stay_id=stay.id,
        visa_type="DN1",
        entry_date=datetime.date(2026, 1, 10),
        expiry_date=datetime.date(2026, 12, 31),
    )
    db_session.add(visa)

    wp = WorkPermit(
        employee_id=emp.id,
        permit_number="GPLD-9999",
        issue_date=datetime.date(2026, 2, 1),
        valid_from=datetime.date(2026, 2, 1),
        valid_to=datetime.date(2028, 2, 1),
        issue_type="CAP_MOI",
    )
    db_session.add(wp)

    contract = Contract(
        employee_id=emp.id,
        contract_number="HDLD-2026-001",
        contract_type="CẤP MỚI",
        start_date=datetime.date(2026, 1, 10),
        end_date=datetime.date(2028, 1, 10),
    )
    db_session.add(contract)
    db_session.flush()

    # Add attachments for Passport and Work Permit
    att_passport = DocumentAttachment(
        entity_type="PASSPORT",
        entity_id=emp.id,
        file_name="passport.pdf",
        file_path="/tmp/passport.pdf",
        file_size=1024,
        mime_type="application/pdf",
    )
    att_wp = DocumentAttachment(
        entity_type="WORK_PERMIT",
        entity_id=wp.id,
        file_name="gpld.pdf",
        file_path="/tmp/gpld.pdf",
        file_size=2048,
        mime_type="application/pdf",
    )
    db_session.add_all([att_passport, att_wp])
    db_session.commit()

    # Generate Excel
    excel_bytes = generate_legal_profile_excel(db_session)
    assert isinstance(excel_bytes, io.BytesIO)

    wb = load_workbook(excel_bytes)
    assert "Danh sách Nhân sự" in wb.sheetnames
    assert "Lịch sử Giấy tờ" in wb.sheetnames

    sheet1 = wb["Danh sách Nhân sự"]
    assert sheet1.cell(row=1, column=1).value == "STT"
    assert sheet1.cell(row=2, column=2).value == "NY1001"
    assert sheet1.cell(row=2, column=3).value == "WANG LEI"
    assert sheet1.cell(row=2, column=19).value == "Cố định"
    assert sheet1.cell(row=2, column=26).value == "GPLD-9999"
    assert sheet1.cell(row=2, column=27).value == "CAP_MOI"
    assert sheet1.cell(row=2, column=30).value == "HDLD-2026-001"
    assert sheet1.cell(row=2, column=31).value == "CẤP MỚI"

    # Verify hyperlinks on Sheet 1
    assert sheet1.cell(row=2, column=11).hyperlink.target == "Giay_To_Dinh_Kem/NY1001_WANG_LEI/passport.pdf"
    assert sheet1.cell(row=2, column=26).hyperlink.target == "Giay_To_Dinh_Kem/NY1001_WANG_LEI/gpld.pdf"

    sheet2 = wb["Lịch sử Giấy tờ"]
    assert sheet2.cell(row=1, column=1).value == "STT"
    # Verify hyperlink on Sheet 2 for WorkPermit row (row 3: header=1, visa=2, gpld=3, contract=4)
    assert sheet2.cell(row=3, column=5).hyperlink.target == "Giay_To_Dinh_Kem/NY1001_WANG_LEI/gpld.pdf"


def test_generate_legal_profile_excel_with_empty_fields_and_fallback_attachment(db_session):
    emp = ForeignEmployee(
        employee_code="NY1005",
        name_latin="LI XIANG",
        gender="Nam",
        passport_number=None,  # Empty passport number
    )
    db_session.add(emp)
    db_session.flush()

    stay = Stay(
        employee_id=emp.id,
        accommodation_type="KTX",
        stay_type="CO_DINH",
        has_meals=True,
    )
    db_session.add(stay)
    db_session.flush()

    visa1 = Visa(stay_id=stay.id, visa_type=None, expiry_date=datetime.date(2025, 1, 1))
    visa2_latest = Visa(stay_id=stay.id, visa_type=None, expiry_date=datetime.date(2026, 1, 1))
    db_session.add_all([visa1, visa2_latest])
    db_session.flush()

    # Add passport attachment to employee, and visa attachment to visa1 (older visa)
    att_passport = DocumentAttachment(
        entity_type="PASSPORT",
        entity_id=emp.id,
        file_name="passport_scan.jpg",
        file_path="/tmp/passport_scan.jpg",
        file_size=1024,
        mime_type="image/jpeg",
    )
    att_visa = DocumentAttachment(
        entity_type="VISA",
        entity_id=visa1.id,
        file_name="visa_scan.jpg",
        file_path="/tmp/visa_scan.jpg",
        file_size=2048,
        mime_type="image/jpeg",
    )
    db_session.add_all([att_passport, att_visa])
    db_session.commit()

    excel_bytes = generate_legal_profile_excel(db_session)
    wb = load_workbook(excel_bytes)
    sheet1 = wb["Danh sách Nhân sự"]

    # Row 2 (NY1005): column 11 (passport) was empty, should now be "Xem file" with hyperlink
    assert sheet1.cell(row=2, column=11).value == "Xem file"
    assert sheet1.cell(row=2, column=11).hyperlink.target == "Giay_To_Dinh_Kem/NY1005_LI_XIANG/passport_scan.jpg"

    # Column 21 (visa) was empty and attachment was on older visa1, should now be "Xem file" with hyperlink
    assert sheet1.cell(row=2, column=21).value == "Xem file"
    assert sheet1.cell(row=2, column=21).hyperlink.target == "Giay_To_Dinh_Kem/NY1005_LI_XIANG/visa_scan.jpg"


def test_generate_legal_profile_excel_multiple_attachments_folder_link(db_session):
    emp = ForeignEmployee(
        employee_code="NY1006",
        name_latin="CHEN WEI",
        gender="Nam",
    )
    db_session.add(emp)
    db_session.flush()

    # Add 2 attachments to Passport entity
    att1 = DocumentAttachment(
        entity_type="PASSPORT",
        entity_id=emp.id,
        file_name="passport_page1.jpg",
        file_path="/tmp/p1.jpg",
        file_size=100,
        mime_type="image/jpeg",
    )
    att2 = DocumentAttachment(
        entity_type="PASSPORT",
        entity_id=emp.id,
        file_name="passport_page2.jpg",
        file_path="/tmp/p2.jpg",
        file_size=100,
        mime_type="image/jpeg",
    )
    db_session.add_all([att1, att2])
    db_session.commit()

    excel_bytes = generate_legal_profile_excel(db_session)
    wb = load_workbook(excel_bytes)
    sheet1 = wb["Danh sách Nhân sự"]

    # Since 2 passport files exist, cell text should be "Xem thư mục (2 file)" and hyperlink target should be the folder
    assert sheet1.cell(row=2, column=11).value == "Xem thư mục (2 file)"
    assert sheet1.cell(row=2, column=11).hyperlink.target == "Giay_To_Dinh_Kem/NY1006_CHEN_WEI"








def test_generate_presence_accommodation_excel(db_session):
    room = Room(room_number="1601", notes="KTX 1601")
    db_session.add(room)
    db_session.flush()

    emp = ForeignEmployee(
        employee_code="NY1002",
        name_latin="ZHANG WEI",
        gender="Nam",
        entry_date=datetime.date(2026, 1, 1),
    )
    db_session.add(emp)
    db_session.flush()

    stay = Stay(
        employee_id=emp.id,
        accommodation_type="KTX",
        room_id=room.id,
        bed_location="Giường A",
        stay_type="CO_DINH",
        has_meals=True,
        start_date=datetime.date(2026, 1, 1),
    )
    db_session.add(stay)
    db_session.commit()

    excel_bytes = generate_presence_accommodation_excel(db_session)
    assert isinstance(excel_bytes, io.BytesIO)

    wb = load_workbook(excel_bytes)
    assert "Sơ đồ KTX & Khách sạn" in wb.sheetnames
    assert "Chưa xếp chỗ ở" in wb.sheetnames
    assert "Đã về nước" in wb.sheetnames


def test_create_report_zip_package():
    excel_io = io.BytesIO(b"Dummy excel content")
    attachments = [
        {
            "employee_code": "NY1001",
            "employee_name": "WANG LEI",
            "file_name": "passport.jpg",
            "content": b"dummy image data",
        }
    ]
    zip_io = create_report_zip_package(excel_io, "Bao_Cao.xlsx", attachments)
    assert isinstance(zip_io, io.BytesIO)
    assert len(zip_io.getvalue()) > 0


def test_export_api_endpoints(client):
    res = client.get("/api/hr-foreign/exports/legal-profile")
    assert res.status_code == 200
    assert res.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    res_zip = client.get("/api/hr-foreign/exports/legal-profile?include_attachments=true")
    assert res_zip.status_code == 200
    assert res_zip.headers["content-type"] == "application/zip"

    res_presence = client.get("/api/hr-foreign/exports/presence-accommodation")
    assert res_presence.status_code == 200
    assert res_presence.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def test_export_legal_profile_zip_with_attachments(client):
    import zipfile
    emp = client.post(
        "/api/hr-foreign/employees",
        json={"name_latin": "WANG LEI", "gender": "Nam", "employee_code": "NV001"},
    ).json()

    files = {"file": ("ho_chieu_wang.pdf", io.BytesIO(b"Fake PDF content"), "application/pdf")}
    data = {"entity_type": "PASSPORT", "entity_id": str(emp["id"])}
    client.post("/api/hr-foreign/attachments", data=data, files=files)

    res_zip = client.get("/api/hr-foreign/exports/legal-profile?include_attachments=true")
    assert res_zip.status_code == 200

    zip_bytes = io.BytesIO(res_zip.content)
    with zipfile.ZipFile(zip_bytes, "r") as zf:
        namelist = zf.namelist()
        assert "Bao_Cao_Ho_So_Phap_Ly.xlsx" in namelist
        assert len(namelist) > 1, f"Expected attachments in ZIP, but got only: {namelist}"


def test_legal_profile_report_excludes_janitors(client):
    from openpyxl import load_workbook
    # Create 1 foreign employee and 1 janitor
    client.post("/api/hr-foreign/employees", json={"name_latin": "EXPAT 1", "gender": "Nam", "employee_type": "FOREIGN_EMPLOYEE"})
    client.post("/api/hr-foreign/employees", json={"name_latin": "JANITOR 1", "gender": "Nữ", "employee_type": "JANITORIAL"})

    res = client.get("/api/hr-foreign/exports/legal-profile")
    assert res.status_code == 200

    excel_bytes = io.BytesIO(res.content)
    wb = load_workbook(excel_bytes)
    ws = wb["Danh sách Nhân sự"]

    names_in_excel = [ws.cell(row=r, column=3).value for r in range(2, ws.max_row + 1)]
    assert "EXPAT 1" in names_in_excel
    assert "JANITOR 1" not in names_in_excel


def test_generate_meal_expense_excel(db_session):
    from features.hr_foreign.excel_exporter import generate_meal_expense_excel
    excel_bytes = generate_meal_expense_excel(db_session, start_date=datetime.date(2026, 8, 1), end_date=datetime.date(2026, 8, 5))
    assert isinstance(excel_bytes, io.BytesIO)

    wb = load_workbook(excel_bytes)
    assert len(wb.sheetnames) == 1
    assert "01.08-05.08" in wb.sheetnames


def test_export_meal_expense_endpoint(client):
    res = client.get("/api/hr-foreign/exports/meal-expense?start_date=2026-08-01&end_date=2026-08-05")
    assert res.status_code == 200
    assert res.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

