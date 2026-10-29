from __future__ import annotations

import datetime as dt
import io
from collections.abc import Generator

import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session, sessionmaker

from core.database import get_db
from features.hr_foreign.models import ForeignEmployee, Room, Stay, TravelRecord
from tests.test_meal_reporting_public import public_client


def test_meal_report_expense_and_excel_follow_current_locks(
    public_client: tuple[TestClient, sessionmaker[Session]], monkeypatch: pytest.MonkeyPatch,
) -> None:
    from features.hr_foreign.meal_reporting import rules
    from main import create_app as create_hr_app

    clock = [dt.datetime(2026, 11, 2, 16, 0, tzinfo=rules.VIETNAM_TZ)]
    monkeypatch.setattr(rules, "now_vietnam", lambda: clock[0])
    portal, session_factory = public_client
    with session_factory.begin() as db:
        room = Room(room_number="1601")
        db.add(room)
        db.flush()
        ids: list[int] = []
        stays: list[int] = []
        for name in ("ONE MEAL", "TWO MEALS", "NO REPORT"):
            employee = ForeignEmployee(name_latin=name, gender="Nam")
            db.add(employee)
            db.flush()
            stay = Stay(employee_id=employee.id, room_id=room.id, accommodation_type="KTX",
                        has_meals=True, start_date=dt.date(2026, 10, 1))
            db.add(stay)
            db.flush()
            db.add(TravelRecord(employee_id=employee.id, entry_date=dt.date(2026, 10, 1)))
            ids.append(employee.id)
            stays.append(stay.id)

    for employee_id, date, breakfast, dinner in (
        (ids[0], "2026-11-03", True, False),
        (ids[1], "2026-11-03", True, True),
        (ids[0], "2026-11-04", False, True),
        (ids[1], "2026-11-04", True, True),
        (ids[1], "2026-11-05", True, True),
    ):
        result = portal.post("/api/meal-reporting/reports", json={
            "employee_id": employee_id, "meal_date": date, "breakfast": breakfast, "dinner": dinner,
        })
        assert result.status_code == 200, result.text

    def override_get_db() -> Generator[Session, None, None]:
        with session_factory() as session:
            yield session

    hr_app = create_hr_app()
    hr_app.dependency_overrides[get_db] = override_get_db
    clock[0] = dt.datetime(2026, 11, 7, 17, 0, tzinfo=rules.VIETNAM_TZ)
    with TestClient(hr_app) as hr:
        assert hr.post("/api/hr-foreign/event-days", json={
            "event_date": "2026-11-04", "event_type": "PRESIDENT_VISIT",
        }).status_code == 201
        assert hr.post(f"/api/hr-foreign/stays/{stays[0]}/meal-absences", json={
            "absence_date": "2026-11-03", "meal_type": "BREAKFAST",
        }).status_code == 201
        assert hr.post(f"/api/hr-foreign/stays/{stays[2]}/meal-extras", json={
            "extra_date": "2026-11-03", "meal_type": "ALL_DAY",
        }).status_code == 201

        def lock(date: str, session: str, suggested: int, final: int, price: int, notes: str | None = None) -> None:
            response = hr.post("/api/hr-foreign/meals/lock", json={
                "lock_date": date, "meal_session": session, "calculated_meal_count": suggested,
                "final_meal_count": final, "locked_price_per_meal": price, "notes": notes,
            })
            assert response.status_code == 201, response.text

        lock("2026-11-03", "BREAKFAST", 2, 2, 30000)
        lock("2026-11-03", "BREAKFAST", 2, 3, 30000, "Khách ăn thêm")
        lock("2026-11-03", "DINNER", 1, 1, 40000)
        lock("2026-11-03", "LUNCH", 1, 1, 25000)
        lock("2026-11-04", "BREAKFAST", 1, 1, 50000)
        lock("2026-11-04", "DINNER", 2, 2, 70000)

        period = "start_date=2026-11-03&end_date=2026-11-08"
        report_response = hr.get(f"/api/hr-foreign/reports/meal-expenses?{period}")
        validation = hr.get(f"/api/hr-foreign/reports/meal-expenses/validate-locks?{period}")
        excel_response = hr.get(f"/api/hr-foreign/exports/meal-expense?{period}")

    assert report_response.status_code == 200, report_response.text
    report = report_response.json()
    assert report["total_meals"] == 10
    assert report["total_expense"] == 415000
    assert report["total_fruit_expense"] == 300000
    assert report["grand_total_expense"] == 715000
    assert report["total_janitor_meals"] == 1
    assert report["total_janitor_expense"] == 25000
    assert next(item for item in report["items"] if item["employee_name"] == "NO REPORT")["absent_days"] == 0
    assert any(item["date"] == "2026-11-05" for item in validation.json()["missing_dates"])
    assert not any(item["date"] == "2026-11-08" for item in validation.json()["missing_dates"])

    assert excel_response.status_code == 200
    sheet = openpyxl.load_workbook(io.BytesIO(excel_response.content)).active
    assert [(sheet.cell(row, 4).value, sheet.cell(row, 7).value) for row in (8, 9, 10)] == [(3, 1), (1, 2), (1, 1)]
    assert [(sheet.cell(row, 5).value, sheet.cell(row, 8).value) for row in (8, 9)] == [(30000, 40000), (50000, 70000)]
    assert sheet.cell(8, 11).value == 1
    assert sheet.cell(13, 4).value == "-"
