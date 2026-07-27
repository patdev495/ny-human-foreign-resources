from __future__ import annotations

import datetime
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from main import app
from core.database import Base, get_db
from features.hr_foreign.models import MealPriceConfig
from features.hr_foreign.services.meal_price_service import seed_default_meal_prices


@pytest.fixture
def client():
    engine = create_engine(
        "sqlite:///:memory:",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(bind=engine)
    TestingSessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)

    def override_get_db():
        db = TestingSessionLocal()
        try:
            yield db
        finally:
            db.close()

    app.dependency_overrides[get_db] = override_get_db
    with TestClient(app) as c:
        yield c
    app.dependency_overrides.clear()


def test_meal_price_config_includes_fruit_allowance(client: TestClient) -> None:
    """Cycle 1: Test MealPriceConfig schema & seed includes fruit_allowance_price (default 60000)."""
    resp = client.get("/api/hr-foreign/meal-price-configs")
    assert resp.status_code == 200
    configs = resp.json()
    assert len(configs) >= 1
    normal_cfg = next(c for c in configs if c["day_type"] == "NORMAL")
    assert "fruit_allowance_price" in normal_cfg
    assert normal_cfg["fruit_allowance_price"] == 60000.0


def test_employee_workplace_location(client: TestClient) -> None:
    """Cycle 2: Test workplace_location on employee schema & default value (DORMITORY)."""
    # Create employee with default DORMITORY workplace_location
    resp1 = client.post(
        "/api/hr-foreign/employees",
        json={
            "name_latin": "NGUYEN THI HOA",
            "gender": "Nu",
            "nationality": "Viet Nam",
            "role": "Tap vu",
        },
    )
    assert resp1.status_code == 201
    emp1 = resp1.json()
    assert emp1["workplace_location"] == "DORMITORY"

    # Create employee with COMPANY workplace_location
    resp2 = client.post(
        "/api/hr-foreign/employees",
        json={
            "name_latin": "TRAN THI MAI",
            "gender": "Nu",
            "nationality": "Viet Nam",
            "role": "Tap vu",
            "workplace_location": "COMPANY",
        },
    )
    assert resp2.status_code == 201
    emp2 = resp2.json()
    assert emp2["workplace_location"] == "COMPANY"


def test_validate_meal_session_locks_period(client: TestClient) -> None:
    """Cycle 3: Test validation of meal session locks for a date range (skipping Sundays)."""
    resp = client.get("/api/hr-foreign/reports/meal-expenses/validate-locks?start_date=2026-08-03&end_date=2026-08-09")
    assert resp.status_code == 200
    data = resp.json()
    # 2026-08-03 is Mon, 2026-08-09 is Sun.
    # Sundays (2026-08-09) should NOT be in missing_dates!
    missing_dates = data.get("missing_dates", [])
    unclosed_days = [item["date"] for item in missing_dates]
    assert "2026-08-09" not in unclosed_days  # Sunday skipped
    assert "2026-08-03" in unclosed_days     # Monday missing lock


def test_generate_daily_meal_expense_excel(client: TestClient) -> None:
    """Cycle 4: Test Excel generation of 1-sheet daily meal expense report matching company template."""
    import openpyxl
    import io
    from features.hr_foreign.exporters.meal_expense_exporter import generate_meal_expense_excel
    from core.database import get_db

    # Create dummy request via export endpoint
    resp = client.get("/api/hr-foreign/exports/meal-expense?start_date=2026-08-03&end_date=2026-08-09")
    assert resp.status_code == 200
    assert resp.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    wb = openpyxl.load_workbook(io.BytesIO(resp.content), data_only=False)
    # 1. Must be exactly 1 sheet
    assert len(wb.sheetnames) == 1
    sheet = wb.active
    assert "03.08" in sheet.title

    # 2. Check title header
    assert "DANH SÁCH THANH TOÁN TIỀN ĂN" in str(sheet.cell(4, 1).value)

    # 3. Check columns in row 6/7
    assert sheet.cell(6, 1).value == "Stt"
    assert "Buổi sáng" in str(sheet.cell(6, 3).value)
    assert "Buổi tối" in str(sheet.cell(6, 6).value)
    assert "hoa quả" in str(sheet.cell(6, 9).value)
    assert "Suất ăn tạp vụ" in str(sheet.cell(6, 10).value)
    assert "Tổng tiền" in str(sheet.cell(6, 13).value)
