from __future__ import annotations

import datetime
import io
import pytest
from openpyxl import load_workbook
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from main import app
from core.database import Base, get_db
from features.hr_foreign.models import ForeignEmployee
from features.hr_foreign.exporters.janitor_payroll_exporter import generate_janitor_payroll_excel


@pytest.fixture
def db_session():
    engine = create_engine(
        "sqlite:///:memory:",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(bind=engine)
    TestingSessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)
    session = TestingSessionLocal()
    try:
        yield session
    finally:
        session.close()


@pytest.fixture
def client(db_session):
    def override_get_db():
        yield db_session

    app.dependency_overrides[get_db] = override_get_db
    with TestClient(app) as c:
        yield c
    app.dependency_overrides.clear()


def test_generate_janitor_payroll_excel(db_session):
    # Setup test data: 2 janitors
    j1 = ForeignEmployee(
        employee_code="NYV2007077",
        name_latin="Đỗ Thị Lan",
        name_chinese="Đỗ Thị Lan",
        gender="Nữ",
        nationality="Việt Nam",
        department="Nhân Sự/管理课",
        role="Tạp vụ CN09",
        work_type="CO_DINH",
        workplace_location="CN09",
        salary=8500000.0,
        salary_unit="MONTH",
        employee_type="JANITORIAL",
    )
    j2 = ForeignEmployee(
        employee_code="NYV2007082",
        name_latin="Nguyễn Thị Thứ",
        name_chinese="Nguyễn Thị Thứ",
        gender="Nữ",
        nationality="Việt Nam",
        department="Nhân Sự/管理课",
        role="Tạp vụ KTX",
        work_type="CO_DINH",
        workplace_location="DORMITORY",
        salary=70000.0,
        salary_unit="DAY",
        employee_type="JANITORIAL",
    )

    # Foreign employee (should NOT be included)
    fe = ForeignEmployee(
        employee_code="NYV1001",
        name_latin="WANG WEI",
        name_chinese="王伟",
        gender="Nam",
        nationality="Trung Quốc",
        department="Kỹ thuật",
        role="Chuyên gia",
        work_type="CO_DINH",
        workplace_location="DORMITORY",
        salary=20000000.0,
        salary_unit="MONTH",
        employee_type="FOREIGN",
    )
    db_session.add_all([j1, j2, fe])
    db_session.commit()

    start_date = datetime.date(2026, 7, 1)
    end_date = datetime.date(2026, 7, 31)

    excel_bytes = generate_janitor_payroll_excel(db_session, start_date, end_date)
    assert isinstance(excel_bytes, io.BytesIO)

    wb = load_workbook(excel_bytes, data_only=False)
    sheet_name = f"{start_date.strftime('%d.%m')}-{end_date.strftime('%d.%m')}"
    assert sheet_name in wb.sheetnames

    ws = wb[sheet_name]
    # Check title row contains 'BẢNG CÔNG-LƯƠNG TẠP VỤ'
    assert "BẢNG CÔNG-LƯƠNG TẠP VỤ" in str(ws.cell(row=5, column=1).value or "")

    # Row 8 should be j1 attendance row
    assert ws.cell(row=8, column=2).value == "ngày công"
    assert ws.cell(row=8, column=3).value == "NYV2007077"
    assert ws.cell(row=8, column=4).value == "Đỗ Thị Lan"
    assert ws.cell(row=8, column=8).value == "CN09"  # Nơi làm việc column

    # Row 9 should be j1 overtime row
    assert ws.cell(row=9, column=2).value == "tăng ca\n" or ws.cell(row=9, column=2).value == "tăng ca"

    # Row 10 should be j2 attendance row
    assert ws.cell(row=10, column=3).value == "NYV2007082"
    assert ws.cell(row=10, column=4).value == "Nguyễn Thị Thứ"
    assert ws.cell(row=10, column=8).value == "DORMITORY"

    # Check Sunday logic: July 2026 5th, 12th, 19th, 26th are Sundays
    # 01/07 is Wed (Col K / Col 11). 05/07 is Sun (Col O / Col 15).
    # Sunday cell should be None / empty
    assert ws.cell(row=8, column=11).value == "N"
    assert ws.cell(row=8, column=15).value is None or ws.cell(row=8, column=15).value == ""

    # Check formula in attendance total column (col AR / 44): =COUNTIFS(...)
    ar8_val = str(ws.cell(row=8, column=44).value or "")
    assert "COUNTIFS" in ar8_val

    # Check formula in salary total column (col AU / 47): =IF(...) for MONTH unit, =AS10*AR10 for DAY unit
    au8_val = str(ws.cell(row=8, column=47).value or "")
    assert "IF" in au8_val or "AS8" in au8_val

    au10_val = str(ws.cell(row=10, column=47).value or "")
    assert "AS10" in au10_val and "AR10" in au10_val


def test_export_janitor_payroll_endpoint(client, db_session):
    j1 = ForeignEmployee(
        employee_code="NYV2007077",
        name_latin="DO THI LAN",
        name_chinese="Đỗ Thị Lan",
        gender="Nữ",
        nationality="Việt Nam",
        department="Nhân Sự/管理课",
        role="Tạp vụ CN09",
        work_type="CO_DINH",
        workplace_location="CN09",
        salary=8500000.0,
        salary_unit="MONTH",
        employee_type="JANITORIAL",
    )
    db_session.add(j1)
    db_session.commit()

    res = client.get(
        "/api/hr-foreign/reports/janitor-payroll/export",
        params={"start_date": "2026-07-01", "end_date": "2026-07-31"},
    )
    assert res.status_code == 200
    assert "spreadsheetml.sheet" in res.headers["content-type"]
    assert "BCC_LUONG_TAP_VU_20260701_20260731.xlsx" in res.headers.get("content-disposition", "")

