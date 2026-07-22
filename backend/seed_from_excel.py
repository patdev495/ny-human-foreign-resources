# -*- coding: utf-8 -*-
"""
Seed script: Đọc 18 nhân viên từ '1. DANH SÁCH NNN 2026.xlsx' làm nguồn chính.
             Đọc KTX file để bổ sung lịch sử chỗ ở (phòng, giường, ngày vào).
             Xoá toàn bộ dữ liệu cũ trước khi seed.
"""
from __future__ import annotations

import datetime
import glob
import os
import re
import sys

import openpyxl
from sqlalchemy.orm import Session

sys.stdout.reconfigure(encoding="utf-8")

from core.database import Base, SessionLocal, engine
from features.hr_foreign import models, service


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def parse_date(val: object) -> datetime.date | None:
    if isinstance(val, datetime.datetime):
        return val.date()
    if isinstance(val, datetime.date):
        return val
    if isinstance(val, (int, float)):
        try:
            return datetime.date(1899, 12, 30) + datetime.timedelta(days=int(val))
        except Exception:
            return None
    if isinstance(val, str):
        val = val.strip()
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%Y/%m/%d", "%d-%m-%Y"):
            try:
                return datetime.datetime.strptime(val, fmt).date()
            except ValueError:
                pass
    return None


def clean_str(val: object) -> str | None:
    if val is None:
        return None
    s = str(val).strip().replace("\t", "")
    return s if s else None


def find_data_file(pattern: str) -> str:
    """Find a file in the data directory by a keyword pattern."""
    for base in [
        os.path.join("..", "data"),
        os.path.join("data"),
        os.path.join("..", "..", "data"),
    ]:
        files = glob.glob(os.path.join(base, "*.xlsx"))
        matches = [f for f in files if pattern.upper() in os.path.basename(f).upper()]
        if matches:
            return matches[0]
    raise FileNotFoundError(f"Cannot find data file matching '{pattern}'")


# ---------------------------------------------------------------------------
# Seed from DANH SÁCH NNN 2026
# ---------------------------------------------------------------------------

def seed_from_danh_sach(db: Session, path: str) -> dict[str, int]:
    """
    Seed ForeignEmployee + WorkPermit + Visa from '1. DANH SÁCH NNN 2026.xlsx'.
    Returns mapping: name_latin (upper) -> employee_id
    """
    print(f"  Loading DANH SACH: {path}")
    wb = openpyxl.load_workbook(path, data_only=True)
    sh = wb[wb.sheetnames[0]]

    name_to_id: dict[str, int] = {}
    emp_count = 0
    wp_count = 0
    visa_count = 0

    for row_idx, row in enumerate(sh.iter_rows(values_only=True), start=1):
        # Header at row 4, data from row 5
        if row_idx < 5:
            continue

        name_raw = clean_str(row[3])
        if not name_raw:
            continue

        # --- Employee basic info ---
        employee_code = clean_str(row[2])   # Mã số (NY...)
        gender_raw = clean_str(row[5])
        nationality = clean_str(row[6])
        passport_raw = clean_str(row[7])
        department = clean_str(row[8])       # Vị trí công việc (category)
        role = clean_str(row[9])             # Chức danh chi tiết
        dob = parse_date(row[4])

        gender = "Nữ" if gender_raw and "nữ" in gender_raw.lower() else "Nam"
        passport_number = str(passport_raw) if passport_raw else None

        emp = models.ForeignEmployee(
            employee_code=employee_code,
            name_latin=name_raw,
            gender=gender,
            nationality=nationality,
            date_of_birth=dob,
            passport_number=passport_number,
            passport_expiry=None,
            department=department,
            role=role,
        )
        db.add(emp)
        db.flush()
        emp_count += 1
        name_to_id[name_raw.upper()] = emp.id

        # --- Work Permit (GPLĐ) ---
        permit_number = clean_str(row[1])    # Số GPLĐ
        gpld_issue = parse_date(row[12])     # Ngày cấp
        gpld_from = parse_date(row[13])      # Từ ngày
        gpld_to = parse_date(row[14])        # Đến ngày
        issue_type = clean_str(row[15])      # Cấp mới / Cấp lại / Gia hạn

        if permit_number or gpld_from or gpld_to:
            wp = models.WorkPermit(
                employee_id=emp.id,
                permit_number=permit_number,
                issue_date=gpld_issue,
                valid_from=gpld_from,
                valid_to=gpld_to,
                issue_type=issue_type,
            )
            db.add(wp)
            wp_count += 1

        # --- Create a default active Stay (no room) to hold Visa records ---
        stay = models.Stay(
            employee_id=emp.id,
            accommodation_type="KTX",
            room_id=None,
            bed_location=None,
            stay_type="CO_DINH",
            has_meals=True,
            start_date=gpld_from,
            end_date=None,
        )
        db.add(stay)
        db.flush()

        # --- Visa lần 1 (col 17-19) ---
        v1_type = clean_str(row[17])
        v1_from = parse_date(row[18])
        v1_to = parse_date(row[19])
        if v1_type or v1_from or v1_to:
            db.add(models.Visa(
                stay_id=stay.id,
                visa_type=v1_type,
                entry_date=v1_from,
                expiry_date=v1_to,
            ))
            visa_count += 1

        # --- Visa lần 2 (col 20-22) ---
        v2_type = clean_str(row[20])
        v2_from = parse_date(row[21])
        v2_to = parse_date(row[22])
        if v2_type or v2_from or v2_to:
            db.add(models.Visa(
                stay_id=stay.id,
                visa_type=v2_type,
                entry_date=v2_from,
                expiry_date=v2_to,
            ))
            visa_count += 1

        # --- Contract (Hợp đồng) (col 10-11) ---
        contract_start = parse_date(row[10])
        contract_end = parse_date(row[11])
        contract_status = clean_str(row[16])  # trạng thái / ghi chú
        
        contract_type = "CẤP MỚI"
        if issue_type:
            val_upper = issue_type.upper()
            if "GIA HẠN" in val_upper:
                contract_type = "GIA HẠN"
            elif "CẤP LẠI" in val_upper or "CẤP ĐỔI" in val_upper:
                contract_type = "CẤP LẠI"
            elif "CHẤM DỨT" in val_upper:
                contract_type = "CHẤM DỨT SỚM"
            elif "CẤP MỚI" in val_upper or "MIỄN" in val_upper:
                contract_type = "CẤP MỚI"
            else:
                contract_type = issue_type.upper()

        if contract_start or contract_end:
            db.add(models.Contract(
                employee_id=emp.id,
                contract_type=contract_type,
                start_date=contract_start,
                end_date=contract_end,
                notes=contract_status,
            ))

    db.flush()
    print(f"    => {emp_count} employees, {wp_count} work permits, {visa_count} visas seeded.")
    return name_to_id


# ---------------------------------------------------------------------------
# Enrich Stay data from KTX file
# ---------------------------------------------------------------------------

def enrich_from_ktx(db: Session, ktx_path: str, name_to_id: dict[str, int]) -> None:
    """
    Read KTX Excel, find rows matching the 18 employees,
    create/update Room + Stay with room & bed info.
    """
    print(f"  Loading KTX: {ktx_path}")
    wb = openpyxl.load_workbook(ktx_path, data_only=True)
    sh = wb["2026"]

    rooms_map: dict[str, models.Room] = {}
    current_room: str | None = None
    enriched = 0

    for row_idx, row in enumerate(sh.iter_rows(values_only=True), start=1):
        if row_idx < 7:
            continue

        # Track current room number (merged cells)
        room_val = clean_str(row[0])
        if room_val and re.match(r"^\d{4}$", room_val):
            current_room = room_val

        bed_val = clean_str(row[1])
        name_raw = clean_str(row[2])
        if not name_raw:
            continue

        if "PHÒNG ĂN" in name_raw.upper() or "PHÒNG THỂ DỤC" in name_raw.upper():
            continue

        # Normalise name: strip extra text in parens, strip trailing space
        name_key = re.sub(r"\s*\(.*?\)", "", name_raw).strip().upper()

        emp_id = name_to_id.get(name_key)
        if not emp_id:
            # Try partial match (name may have trailing space in KTX)
            for key, eid in name_to_id.items():
                if key.startswith(name_key) or name_key.startswith(key):
                    emp_id = eid
                    break

        if not emp_id:
            continue  # Not one of our 18 people

        # Ensure Room exists
        if current_room:
            if current_room not in rooms_map:
                room_db = db.query(models.Room).filter(models.Room.room_number == current_room).first()
                if not room_db:
                    room_db = models.Room(room_number=current_room)
                    db.add(room_db)
                    db.flush()
                rooms_map[current_room] = room_db
            room_obj = rooms_map[current_room]
        else:
            room_obj = None

        entry_date = parse_date(row[7])
        notes_raw = clean_str(row[12])

        # Update the default Stay created from DANH SÁCH
        active_stay = (
            db.query(models.Stay)
            .filter(
                models.Stay.employee_id == emp_id,
                models.Stay.end_date.is_(None),
            )
            .first()
        )

        if active_stay and room_obj:
            active_stay.room_id = room_obj.id
            active_stay.bed_location = bed_val
            if entry_date and not active_stay.start_date:
                active_stay.start_date = entry_date
            if notes_raw:
                active_stay.notes = notes_raw
            enriched += 1

    db.flush()
    print(f"    => {enriched} stays enriched with room/bed data. {len(rooms_map)} rooms created.")


# ---------------------------------------------------------------------------
# Main seed
# ---------------------------------------------------------------------------

def seed() -> None:
    print("=" * 60)
    print("Resetting database...")
    Base.metadata.drop_all(bind=engine)
    Base.metadata.create_all(bind=engine)

    db: Session = SessionLocal()
    try:
        # Seed default meal prices
        service.seed_default_meal_prices(db)

        # Step 1: Seed 18 employees from DANH SÁCH
        danh_sach_path = find_data_file("DANH")
        name_to_id = seed_from_danh_sach(db, danh_sach_path)

        # Step 2: Enrich stays from KTX file
        try:
            ktx_path = find_data_file("KTX")
            enrich_from_ktx(db, ktx_path, name_to_id)
        except FileNotFoundError as e:
            print(f"  Warning: {e} — skipping KTX enrichment.")

        db.commit()

        # Summary
        emp_count = db.query(models.ForeignEmployee).count()
        wp_count = db.query(models.WorkPermit).count()
        stay_count = db.query(models.Stay).count()
        room_count = db.query(models.Room).count()
        visa_count = db.query(models.Visa).count()

        print("=" * 60)
        print("Seed complete!")
        print(f"  Employees  : {emp_count}")
        print(f"  WorkPermits: {wp_count}")
        print(f"  Stays      : {stay_count}")
        print(f"  Rooms      : {room_count}")
        print(f"  Visas      : {visa_count}")
        print("=" * 60)

    except Exception as e:
        db.rollback()
        print(f"ERROR: {e}")
        raise
    finally:
        db.close()


if __name__ == "__main__":
    seed()
